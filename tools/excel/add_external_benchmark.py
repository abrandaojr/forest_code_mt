"""Add the same external-comparison sheet to every final project workbook."""
from __future__ import annotations

import html
import os
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
FORMULA_FILES = sorted((ROOT / "deliverables/01_excel").glob("*.xlsx"))
STANDARD_FILES = [
    ROOT / "out/table/forest_code_gta_final_mt_20260818.xlsx",
    ROOT / "out/table/forest_code_mt_priority_tables_20260818.xlsx",
    ROOT / "out/table/masson_style_final_tables_figures_20260818.xlsx",
]

ROWS = [
    ["EXTERNAL BENCHMARK — MATO GROSSO FOREST CODE COMPLIANCE"],
    ["Status", "Provisional: project estimates require recomputation after correction of APP partitions and 2000–2008 spatial chronology."],
    ["Metric", "Current project", "OCF 2019", "OCF 2024", "SFB / SICAR", "Project − OCF reference", "Interpretation"],
    ["Properties / registrations", 169533, 116390, 176629, "Monthly administrative panel", "=B4-C4", "Different CAR status and property universes"],
    ["Registered area (Mha)", 75.9624, 64.5009, 62.4903, "Extraction-date dependent", "=B5-C5", "Overlap treatment materially affects area"],
    ["Legal Reserve deficit (Mha)", 7.3685, 2.3, None, "Administrative passives after analysis", "=B6-C6", "Project uses adjusted modeled deficit"],
    ["APP metric (Mha)", 0.1589, 0.505, None, "Administrative passives after analysis", "=B7-C7", "Project is restoration; OCF is deficit"],
    ["Total environmental deficit (Mha)", 7.5274, 2.805, 4.7389, "Consult current panel", "=B8-D8", "Components and legal treatment differ"],
    ["Surplus (Mha)", 16.8287, 8.9, 5.5373, "Consult current panel", "=B9-D9", "Project estimate is materially higher"],
    [],
    ["Comparison protocol", "Recompute", "Harmonize population", "Disaggregate metrics", "Compare municipalities", "Inspect spatial outliers", "Report uncertainty"],
    ["Sources", "OCF Mato Grosso diagnostic (2019)", "OCF Forest Code Thermometer (2024)", "SFB/SICAR Rural Environmental Regularization Panel"],
    ["URLs", "https://observatorioflorestal.org.br/wp-content/uploads/2024/12/diagnostico-mt-v01.pdf", "https://observatorioflorestal.org.br/wp-content/uploads/2025/10/BoletimTermometro_do_Codigo_Florestal_2024.pdf", "https://dados.florestal.gov.br/km/dataset/painel-da-regularizacao-ambiental-rural"],
]


def cell(ref: str, value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return f'<c r="{ref}"><v>{value}</v></c>'
    if isinstance(value, str) and value.startswith("="):
        return f'<c r="{ref}"><f>{html.escape(value[1:])}</f><v>0</v></c>'
    return f'<c r="{ref}" t="inlineStr"><is><t>{html.escape(str(value))}</t></is></c>'


def benchmark_xml() -> str:
    out = ['<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><cols>']
    for i, width in enumerate([27, 30, 28, 28, 34, 25, 48], 1):
        out.append(f'<col min="{i}" max="{i}" width="{width}" customWidth="1"/>')
    out.append('</cols><sheetViews><sheetView workbookViewId="0"><pane ySplit="3" state="frozen" topLeftCell="A4"/></sheetView></sheetViews><sheetData>')
    for r, values in enumerate(ROWS, 1):
        out.append(f'<row r="{r}" ht="{34 if r in (1,2,3,11) else 30}" customHeight="1">')
        for c, value in enumerate(values, 1):
            col = chr(64 + c)
            out.append(cell(f"{col}{r}", value))
        out.append('</row>')
    out.append('</sheetData><mergeCells count="2"><mergeCell ref="A1:G1"/><mergeCell ref="B2:G2"/></mergeCells></worksheet>')
    return "".join(out)


def inject_formula_workbook(path: Path) -> None:
    fd, temp_name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent); os.close(fd)
    temp = Path(temp_name)
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode("utf-8").replace('</sheets>', '<sheet name="External Benchmark" sheetId="3" r:id="rId4"/></sheets>')
                data = text.encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                text = data.decode("utf-8").replace('</Relationships>', '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet3.xml"/></Relationships>')
                data = text.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = data.decode("utf-8").replace('</Types>', '<Override PartName="/xl/worksheets/sheet3.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
                data = text.encode("utf-8")
            elif item.filename == "xl/styles.xml":
                data = data.decode("utf-8").replace('name val="Aptos"', 'name val="Arial"').encode("utf-8")
            dst.writestr(item, data)
        dst.writestr("xl/worksheets/sheet3.xml", benchmark_xml())
    temp.replace(path)


def update_standard_workbook(path: Path) -> None:
    wb = load_workbook(path)
    if "External Benchmark" in wb.sheetnames:
        del wb["External Benchmark"]
    ws = wb.create_sheet("External Benchmark", 0)
    for row in ROWS:
        ws.append(row)
    ws.merge_cells("A1:G1"); ws.merge_cells("B2:G2")
    widths = [27, 30, 28, 28, 34, 25, 48]
    for i, width in enumerate(widths, 1): ws.column_dimensions[chr(64+i)].width = width
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=11, color="27362F")
            c.alignment = Alignment(vertical="center", wrap_text=True)
    for c in ws[1]: c.fill=PatternFill("solid",fgColor="3D2A87"); c.font=Font(name="Arial",size=18,bold=True,color="FFFFFF")
    for c in ws[2]: c.fill=PatternFill("solid",fgColor="FCE7EC"); c.font=Font(name="Arial",size=11,bold=True,color="9F2348")
    for c in ws[3]: c.fill=PatternFill("solid",fgColor="147D72"); c.font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
    for c in ws[11]: c.fill=PatternFill("solid",fgColor="F2C94C"); c.font=Font(name="Arial",size=11,bold=True,color="3D2A87")
    for r in range(1,14): ws.row_dimensions[r].height = 34 if r in (1,2,3,11) else 30
    ws.freeze_panes="A4"
    fd, temp_name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent); os.close(fd); temp=Path(temp_name)
    wb.save(temp); wb.close(); temp.replace(path)


if __name__ == "__main__":
    for p in FORMULA_FILES:
        inject_formula_workbook(p); print(p)
    for p in STANDARD_FILES:
        update_standard_workbook(p); print(p)
