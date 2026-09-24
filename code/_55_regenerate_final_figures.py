"""Regenerate only the final 7.5 x 7.5 inch English figure set.

This uses the already-built final parquet outputs, avoiding a full spatial
reprocessing run when only publication graphics need to change.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

import _00_paths as paths
import _10_kernel as kernel
import _40_build_maps as maps
import _50_build_publication_tables as publication
import _56_build_flow_figures as flow_figures
import _80_write_one_pager as one_pager
import _16_audit_spatial_inputs as input_audit


DELIVERY_FIGURES = paths.FIGURES


def regenerate() -> list[Path]:
    paths.set_env()
    publication.apply_map_style()
    fc = pd.read_parquet(
        publication.today_path("forest_code_mt_priority_consolidated_with_secondary", "parquet")
    )
    joined = pd.read_parquet(
        publication.today_path("forest_code_gta_final_mt_with_secondary", "parquet")
    )
    fc = publication.apply_input_order(publication.prep_fc(fc))
    joined = publication.apply_input_order(publication.prep_fc(joined))
    joined["supplier_type"] = pd.Categorical(
        joined["supplier_type"], publication.SUPPLIER_ORDER, ordered=True
    )
    active = fc[fc["car_valid"]].copy()

    lr = publication.table_lr(active, ["size_class"]).rename(columns={"size_class": "Size Class"})
    combined = publication.table_combined(active, ["size_class"]).rename(
        columns={"size_class": "Size Class"}
    )
    mean_deficit = publication.table_mean_deficit(active)

    generated: dict[str, Path] = {}
    generated["fig03_study_area.png"] = publication.save_study_area_map()
    generated["fig04_rl_deficit.png"] = publication.save_stacked_bar(
        lr,
        "Size Class",
        ["To Restore", "To Compensate"],
        "Figure 4: Legal reserve deficit breakdown",
        paths.FIGURES / "Figure_04_LR_deficit_breakdown.png",
    )
    generated["fig05_rl_balance.png"] = publication.save_grouped_bar(
        pd.DataFrame(
            {
                "Biome": ["Forest", "Forest", "Cerrado", "Cerrado"],
                "Metric": ["Adjusted deficit", "Surplus", "Adjusted deficit", "Surplus"],
                "Area (ha)": [
                    active["rl_adj_deficit_forest_ha"].sum(),
                    active["rl_surplus_forest_ha"].sum(),
                    active["rl_adj_deficit_cerrado_ha"].sum(),
                    active["rl_surplus_cerrado_ha"].sum(),
                ],
            }
        ),
        "Biome",
        "Area (ha)",
        "Metric",
        "Figure 5: LR compensation demand and supply by biome",
        paths.FIGURES / "Figure_05_LR_biome_supply_demand.png",
    )
    generated["fig06_compliance.png"] = publication.save_pie_counts(
        pd.Series(
            [
                ((active["rl_adj_deficit_ha"] + active["app_restore_ha"]) == 0).sum(),
                ((active["rl_adj_deficit_ha"] + active["app_restore_ha"]) > 0).sum(),
            ]
        ),
        pd.Series(["Fully compliant", "Any liability"]),
        "Figure 6: Full compliance status",
        paths.FIGURES / "Figure_06_compliance_status.png",
    )
    generated["fig07_restoration.png"] = publication.save_stacked_bar(
        combined,
        "Size Class",
        ["To Restore (ha)", "To Compensate (ha)"],
        "Figure 7: Liability for restoration and compensation",
        paths.FIGURES / "Figure_07_restoration_compensation.png",
    )
    generated["fig08_deficit_by_size.png"] = publication.save_grouped_bar(
        mean_deficit.assign(Metric="Mean total deficit"),
        "Size Class",
        "Total Mean Deficit (ha)",
        "Metric",
        "Figure 8: Mean deficit per property by size class",
        paths.FIGURES / "Figure_08_mean_deficit_size.png",
    )

    municipal_panels = one_pager.build_paper_map_panel(one_pager.build_municipal_maps())
    generated["fig09a_municipal_outcomes_panel_1_of_2.png"] = municipal_panels[0]
    generated["fig09b_municipal_outcomes_panel_2_of_2.png"] = municipal_panels[1]
    generated["fig10_vegetation.png"] = publication.save_vegetation_cover_figure(active)
    generated["fig11_secondary.png"] = publication.save_secondary_impact_figure(active)
    generated["fig12_rule2000.png"] = publication.save_cons2000_overall_figure(active)
    generated["fig13_rule2000_size.png"] = publication.save_cons2000_size_figure(active)
    generated["fig14_suppliers.png"] = publication.save_supplier_subgroups_figure(joined)
    generated["fig15_suppliers_rule2000.png"] = publication.save_supplier_cons2000_figure(joined)

    base_fc = pd.read_parquet(maps.today_path("forest_code_mt_priority_consolidated", "parquet"))
    base_joined = pd.read_parquet(maps.today_path("forest_code_gta_final_mt", "parquet"))
    map_paths = maps.make_maps(maps.build_tables(base_fc, base_joined)["municipal_fc"])
    map_names = {
        "map_app_gross_deficit_ha.png": "map_app_deficit.png",
        "map_rl_adjusted_deficit_ha.png": "map_rl_deficit.png",
        "map_total_deficit_ha.png": "map_total_deficit.png",
    }
    for source in map_paths:
        generated[map_names[source.name]] = source

    input_panels = input_audit.write_outputs(input_audit.load_inventory())
    for panel_number, source in enumerate(input_panels, start=1):
        generated[f"fig16_spatial_inputs_panel_{panel_number:02d}_of_{len(input_panels):02d}.png"] = source

    DELIVERY_FIGURES.mkdir(parents=True, exist_ok=True)
    for obsolete in DELIVERY_FIGURES.glob("fig16*_spatial_inputs_panel_*.png"):
        obsolete.unlink()
    for obsolete_name in (
        "board01_overview.png", "board02_app_rl.png", "board03_regularization.png",
        "fig09_municipal.png",
    ):
        obsolete = DELIVERY_FIGURES / obsolete_name
        if obsolete.exists():
            obsolete.unlink()
    copied = []
    for final_name, source in generated.items():
        if source is None:
            raise RuntimeError(f"Figure source was not generated: {final_name}")
        destination = DELIVERY_FIGURES / final_name
        shutil.copy2(source, destination)
        copied.append(destination)

    flow_figures.style()
    flow_data = flow_figures.load_data()
    copied.extend([
        flow_figures.waterfall(flow_data),
        flow_figures.sankey(flow_data),
        flow_figures.methodology(),
    ])

    workbook = publication.OUT_DIR / f"masson_style_final_tables_figures_{kernel.DATE}.xlsx"
    if workbook.exists():
        wb = load_workbook(workbook)
        if "Figures_Index" in wb.sheetnames:
            del wb["Figures_Index"]
        ws = wb.create_sheet("Figures_Index")
        ws.append(["figure", "path"])
        row = 2
        for figure in copied:
            ws.append([figure.name, str(figure)])
            with PILImage.open(figure) as source_image:
                aspect = source_image.height / source_image.width
            image = XLImage(str(figure))
            image.width = 720
            image.height = int(720 * aspect)
            ws.add_image(image, f"D{row}")
            row += max(24, int(image.height / 20) + 2)
        publication.format_workbook_tables(wb)
        wb.save(workbook)
    return copied


if __name__ == "__main__":
    outputs = regenerate()
    print(f"Regenerated {len(outputs)} final figures for {kernel.DATE}:")
    for output in outputs:
        print(output)
