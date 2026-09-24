from __future__ import annotations

import importlib.util
import os
import re
import sys
import textwrap
from pathlib import Path

import geopandas as gpd
import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import shapely
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

import _00_paths as paths
import _10_kernel as kernel
import _11_forest_code_compliance as compliance
import _12_gta_supply_chain as gta_chain

SCRIPT_DIR = Path(__file__).resolve().parent
paths.set_env()
ROOT = paths.ROOT
DATA_DIR = paths.RAW
OUTPUT_DIR = paths.OUT
RAW_DIR = paths.RAW
PROCESSED_DIR = paths.PROC
OUT_DIR = paths.TABLES
FIG_DIR = paths.FIGURES
FIG_DIR.mkdir(parents=True, exist_ok=True)

def first_existing(candidates: list[Path]) -> Path:
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


MUN_GEOM = paths.MUNICIPALITIES
UF_GEOM = paths.STATES
FIG_SIZE_IN = (10.0, 5.4)
CHART_SIZE_IN = (10.0, 5.4)
MAP_SIZE_IN = (9.0, 6.2)
FIG_DPI = 300
FIG_FONT_SIZE = 12
TITLE_FONT_SIZE = 18
TABLE_BODY_FONT_SIZE = 12
TABLE_HEADER_FONT_SIZE = 14
DISPLAY_CRS = "ESRI:102033"
CRS_LABEL = "South America Albers Equal Area Conic"

MAKE_A_MAP_CANDIDATES = [
    Path(os.environ["MAKE_A_MAP_SRC"]) if "MAKE_A_MAP_SRC" in os.environ else None,
    ROOT / "code",
]
MAKE_A_MAP_CANDIDATES = [candidate for candidate in MAKE_A_MAP_CANDIDATES if candidate is not None]
for candidate in MAKE_A_MAP_CANDIDATES:
    if candidate.exists() and str(candidate) not in sys.path:
        sys.path.insert(0, str(candidate))
        break

try:
    from make_a_map.components._00_legend import LegendItem, draw_legend
    from make_a_map.components._01_scale import draw_scale_bar
    from make_a_map._01_layout import get_layout
    from make_a_map._00_theme import DEFAULT_THEME
except ImportError:  # pragma: no cover - fallback only if reference repo is absent.
    LegendItem = None
    draw_legend = None
    draw_scale_bar = None
    get_layout = None
    DEFAULT_THEME = None


def load_module(name: str, filename: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPT_DIR / filename)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load {filename}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


fcm_gta = load_module("fcm_gta", "_30_build_gta.py")
input_audit = load_module("input_audit", "_16_audit_spatial_inputs.py")


SIZE_ORDER = [
    "Large (>15 MF)",
    "Medium (>4-15 MF)",
    "Small (1-4 MF)",
    "Minifundio (<1 MF)",
    "Not Classified",
]

INPUT_ORDER = kernel.SOURCE_ORDER
INPUT_LABELS = kernel.SOURCE_LABELS
SUPPLIER_ORDER = gta_chain.TYPE_ORDER
SPECIAL_AREA_SPECS = {
    "Conserv. Units": ["unidades_conservacao", "unidades_conservacao_x_car_atp_fnl"],
    "Quilombola": ["quilombolas", "quilombolas_x_car_atp_fnl"],
    "Rural Settl.": ["assentamentos", "assentamentos_x_car_atp_fnl", "assentamentos_intermat", "assentamentos_intermat_x_car_atp_fnl"],
    "Indigenous": ["terras_indigenas", "terras_indigenas_x_car_atp_fnl"],
}

SECONDARY_FOREST = RAW_DIR / "simcar_proxy" / "input_prodes_native_vegetation_2024_plus_sv.shp"
SECONDARY_BY_PROPERTY = OUT_DIR / f"secondary_vegetation_by_property_{kernel.DATE}.parquet"
CAR_CODE_RE = re.compile(r"(MT-\d{7}-[A-Z0-9]+)")


def today_path(name: str, ext: str) -> Path:
    return OUT_DIR / f"{name}_{kernel.DATE}.{ext}"


def num(df: pd.DataFrame, col: str) -> pd.Series:
    if col not in df.columns:
        return pd.Series(0.0, index=df.index)
    return pd.to_numeric(df[col], errors="coerce").fillna(0.0)


def pct(part: pd.Series | float, total: float) -> pd.Series | float:
    if total == 0:
        return 0
    return part / total * 100


def apply_input_order(df: pd.DataFrame, col: str = "input_file_type") -> pd.DataFrame:
    out = df.copy()
    if col in out.columns:
        out[col] = pd.Categorical(out[col].astype(str), INPUT_ORDER, ordered=True)
        out = out.sort_values(col, kind="stable")
    return out


def display_input_labels(series: pd.Series) -> pd.Series:
    return series.astype(str).map(INPUT_LABELS).fillna(series.astype(str))


def map_theme():
    if DEFAULT_THEME is not None:
        return DEFAULT_THEME
    return None


def apply_map_style() -> None:
    theme = map_theme()
    mpl.rcParams.update({
        "font.family": "Arial",
        "font.size": FIG_FONT_SIZE,
        "axes.labelsize": FIG_FONT_SIZE,
        "xtick.labelsize": FIG_FONT_SIZE,
        "ytick.labelsize": FIG_FONT_SIZE,
        "legend.fontsize": FIG_FONT_SIZE,
        "legend.title_fontsize": FIG_FONT_SIZE,
        "axes.titlesize": TITLE_FONT_SIZE,
        "axes.unicode_minus": False,
        "savefig.dpi": FIG_DPI,
        "savefig.facecolor": "white",
        "savefig.bbox": "tight",
        "savefig.pad_inches": 0.08,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
    })


def map_colors() -> dict[str, str]:
    theme = map_theme()
    if theme is None:
        return {
            "paper": "#FFFFFF",
            "ink": "#222222",
            "muted": "#6E6D68",
            "context_fill": "#DDDCD6",
            "context_line": "#B9B8B2",
            "boundary_light": "#FFFFFF",
            "primary": "#1E786A",
            "primary_edge": "#11574D",
            "accent": "#D67C42",
        }
    return {
        "paper": "#FFFFFF",
        "ink": theme.colors.ink,
        "muted": theme.colors.muted_ink,
        "context_fill": theme.colors.context_fill,
        "context_line": theme.colors.context_line,
        "boundary_light": theme.colors.boundary_light,
        "primary": theme.colors.data_primary,
        "primary_edge": theme.colors.data_primary_edge,
        "accent": theme.colors.data_accent,
    }


def map_layout(name: str = "map-plus-metric"):
    if get_layout is not None:
        return get_layout(name)
    raise RuntimeError("make_a_map layout helpers are unavailable")


def wrap_map_text(value: str, width: int, max_lines: int) -> str:
    lines = textwrap.wrap(str(value), width=width, break_long_words=False, break_on_hyphens=False)
    return "\n".join(lines[:max_lines])


def set_map_extent(ax, gdf: gpd.GeoDataFrame, pad_x: float = 0.035, pad_y: float = 0.045) -> None:
    minx, miny, maxx, maxy = gdf.total_bounds
    ax.set_xlim(minx - (maxx - minx) * pad_x, maxx + (maxx - minx) * pad_x)
    ax.set_ylim(miny - (maxy - miny) * pad_y, maxy + (maxy - miny) * pad_y)
    ax.set_aspect("equal")
    ax.set_axis_off()


def add_map_source(fig: plt.Figure, text: str, y: float = 0.033) -> None:
    colors = map_colors()
    fig.text(
        0.055,
        y,
        wrap_map_text(text, 82, 2),
        ha="left",
        va="bottom",
        fontsize=FIG_FONT_SIZE,
        color=colors["muted"],
    )


def compact_number(value: float, _pos=None) -> str:
    if pd.isna(value):
        return ""
    value = float(value)
    av = abs(value)
    if av >= 10:
        return f"{value:,.0f}"
    return f"{value:,.1f}".rstrip("0").rstrip(".")


def compact_percent(value: float, _pos=None) -> str:
    if pd.isna(value):
        return ""
    return f"{float(value):,.0f}%"


def abbreviated_number(value: float, _pos=None) -> str:
    value = float(value)
    if abs(value) >= 1_000_000:
        return f"{value / 1_000_000:.0f}M"
    if abs(value) >= 1_000:
        return f"{value / 1_000:.0f}K"
    return f"{value:.0f}"


def axis_label(text: str) -> str:
    labels = {
        "Properties": "Number of properties",
        "Cattle head": "Cattle head",
        "Area (ha)": "Area (hectares)",
        "Area (million ha)": "Area (million hectares)",
        "Total liability (ha)": "Total liability (hectares)",
        "Mean deficit (ha)": "Mean deficit (hectares)",
    }
    return labels.get(text, text)


def display_header(value: object) -> str:
    text = str(value)
    text = text.replace("(Mha)", "(million hectares)")
    text = text.replace("(ha)", "(hectares)")
    text = text.replace("_ha", " (hectares)")
    text = re.sub(r"\bMha\b", "million hectares", text)
    text = re.sub(r"\bha\b", "hectares", text)
    return text


TABLE_LANGUAGE_RENAMES = {
    "N": "Properties",
    "% N": "% properties",
    "Area (ha)": "Area (hectares)",
    "% Area": "% area",
    "Area_ha": "Area (hectares)",
    "Area overlap (ha)": "Overlap area (hectares)",
    "Property area (ha)": "Property area (hectares)",
    "Compliant n": "Compliant properties",
    "Compliant (n)": "Compliant properties",
    "Liability n": "Affected properties",
    "Any liability (n)": "Affected properties",
    "Properties Net Deficit (n)": "APP-affected properties",
    "LR Def. (ha)": "LR affected area (hectares)",
    "LR Adj. Deficit (ha)": "LR affected area (hectares)",
    "LR Def. (hectares)": "LR affected area (hectares)",
    "APP Liab. (ha)": "APP affected area (hectares)",
    "APP Liab. (hectares)": "APP affected area (hectares)",
    "APP Restore (ha)": "APP affected area (hectares)",
    "APP Net Deficit (ha)": "APP affected area (hectares)",
    "APP Gross Deficit (ha)": "APP gross affected area (hectares)",
    "Total Liab. (ha)": "Total affected area (hectares)",
    "Total Liab. (hectares)": "Total affected area (hectares)",
    "Combined Liability (ha)": "Total affected area (hectares)",
    "Restore (ha)": "Restoration area (hectares)",
    "To Restore (ha)": "Restoration area (hectares)",
    "To Restore": "Restoration area (hectares)",
    "Comp. (ha)": "Compensation area (hectares)",
    "To Compensate (ha)": "Compensation area (hectares)",
    "To Compensate": "Compensation area (hectares)",
    "Required (ha)": "Required LR area (hectares)",
    "Surplus (ha)": "LR surplus area (hectares)",
    "Gross Deficit": "LR gross affected area (hectares)",
    "Adj. Deficit": "LR affected area (hectares)",
    "LR Surplus": "LR surplus area (hectares)",
    "LR Deficit": "LR affected area (hectares)",
    "APP Deficit": "APP affected area (hectares)",
    "Total Deficit": "Total affected area (hectares)",
    "Gross Deficit (ha)": "Gross affected area (hectares)",
    "Adj. Deficit (ha)": "Total affected area (hectares)",
    "Total Mean Deficit (ha)": "Mean hectares/affected property",
    "Mean Net deficit (ha)": "Mean APP hectares/APP-affected property",
    "Baseline LR liability (ha)": "LR with 2000 rule area (hectares)",
    "Without 2000 LR liability (ha)": "LR without 2000 rule area (hectares)",
    "LR delta (ha)": "LR 2000-rule effect area (hectares)",
    "Baseline APP liability (ha)": "APP with 2000 rule area (hectares)",
    "Without 2000 APP liability (ha)": "APP without 2000 rule area (hectares)",
    "APP delta (ha)": "APP 2000-rule effect area (hectares)",
    "Baseline total liability (ha)": "Total with 2000 rule area (hectares)",
    "Without 2000 total liability (ha)": "Total without 2000 rule area (hectares)",
    "Total delta (ha)": "Total 2000-rule effect area (hectares)",
}


def align_table_language(table: pd.DataFrame) -> pd.DataFrame:
    out = table.rename(columns=TABLE_LANGUAGE_RENAMES).copy()
    if "Scenario" in out.columns:
        out["Scenario"] = (
            out["Scenario"]
            .astype(str)
            .replace(
                {
                    "With 2000 rule": "LR with 2000 rule",
                    "Without 2000 rule": "LR without 2000 rule",
                }
            )
        )
    return out


def wrap_axis_labels(labels: list[object] | pd.Series, width: int = 14) -> list[str]:
    return [
        "\n".join(textwrap.wrap(str(label), width=width, break_long_words=False, break_on_hyphens=False)) or str(label)
        for label in labels
    ]


def horizontal_xticks(ax: plt.Axes, labels: list[object] | pd.Series | None = None, width: int = 14) -> None:
    if labels is not None:
        wrapped = wrap_axis_labels(labels, width=width)
        ax.set_xticks(np.arange(len(wrapped)))
        ax.set_xticklabels(wrapped, rotation=0, ha="center", fontsize=FIG_FONT_SIZE)
    else:
        current = [tick.get_text() for tick in ax.get_xticklabels()]
        ax.set_xticklabels(wrap_axis_labels(current, width=width), rotation=0, ha="center", fontsize=FIG_FONT_SIZE)


def add_chart_title(ax: plt.Axes, title: str) -> None:
    """Titles belong in captions and filenames, never inside final figures."""
    return None


def polish_chart(fig: plt.Figure, ax: plt.Axes, ylabel: str | None = None, percent: bool = False) -> None:
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.yaxis.set_major_formatter(FuncFormatter(compact_percent if percent else compact_number))
    if ylabel:
        ax.set_ylabel(axis_label(ylabel), fontsize=FIG_FONT_SIZE, color="#333333")
    ax.tick_params(axis="both", labelsize=FIG_FONT_SIZE, colors="#333333")
    ax.grid(axis="y", color="#DCE4E8", linewidth=0.7)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#777777")
    ax.spines["bottom"].set_color("#777777")


def add_bar_labels(ax: plt.Axes, stacked_totals: np.ndarray | None = None) -> None:
    if stacked_totals is not None:
        threshold = max(stacked_totals, default=0) * 0.03
        for idx, total in enumerate(stacked_totals):
            if total >= threshold and total > 0:
                ax.text(idx, total, compact_number(total), ha="center", va="bottom", fontsize=FIG_FONT_SIZE, color="#333333")
        return
    heights = [bar.get_height() for container in ax.containers for bar in container]
    threshold = max(heights, default=0) * 0.03
    for container in ax.containers:
        labels = [compact_number(bar.get_height()) if bar.get_height() >= threshold and bar.get_height() > 0 else "" for bar in container]
        ax.bar_label(container, labels=labels, padding=2, fontsize=FIG_FONT_SIZE, color="#333333")


def place_legend_right(ax: plt.Axes, *args, **kwargs) -> None:
    """Place a compact legend above the plotting field to maximize data area."""
    handles, _ = ax.get_legend_handles_labels()
    kwargs.update({
        "loc": "lower left",
        "bbox_to_anchor": (0.0, 1.01),
        "frameon": False,
        "fontsize": FIG_FONT_SIZE,
        "ncol": max(1, min(3, len(handles))),
        "borderaxespad": 0.0,
        "handlelength": 1.4,
        "columnspacing": 1.2,
    })
    ax.legend(*args, **kwargs)


def square_chart_layout(fig: plt.Figure) -> None:
    """Apply a consistent landscape editorial layout."""
    fig.subplots_adjust(left=0.11, right=0.98, bottom=0.17, top=0.84)


def mun_code(df: pd.DataFrame) -> pd.Series:
    raw = df["mun_geocodigo"] if "mun_geocodigo" in df.columns else df.get("MUNICIPIO_", pd.Series(pd.NA, index=df.index))
    out = pd.to_numeric(raw, errors="coerce").astype("Int64").astype("string").str.zfill(7)
    return out.mask(out.str.fullmatch(r"0+").fillna(False))


def prep_fc(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if "size_class" not in df.columns:
        df["size_class"] = "Not Classified"
    df["size_class"] = df["size_class"].fillna("Not Classified")
    df["size_class"] = pd.Categorical(df["size_class"], SIZE_ORDER, ordered=True)
    for col in [
        "area_ha_car", "rl_req_total_ha", "rl_gross_deficit_ha", "rl_adj_deficit_ha",
        "rl_restore_ha", "rl_compensate_ha", "rl_surplus_total_ha", "rl_req_forest_ha",
        "rl_req_cerrado_ha", "rl_surplus_forest_ha", "rl_surplus_cerrado_ha",
        "rl_adj_deficit_forest_ha", "rl_adj_deficit_cerrado_ha", "app_req_ha",
        "app_preserved_ha", "app_gross_deficit_ha", "app_restore_ha",
        "app_replant_raw_ha", "app_consolidated_ha", "app_consol_restore_ha",
        "app_restore_auas_ha", "app_cap_ha", "cons_area_2000",
        "auas_post2008", "calc_gross_deficit_total_ha", "calc_deficit_total_ha",
        "secondary_vegetation_ha",
        "rl_exist_forest_with_secondary_ha", "rl_gross_deficit_with_secondary_ha",
        "rl_adj_deficit_with_secondary_ha", "rl_restore_with_secondary_ha",
        "rl_compensate_with_secondary_ha", "rl_surplus_with_secondary_ha",
        "calc_gross_deficit_total_with_secondary_ha",
        "calc_deficit_total_with_secondary_ha", "secondary_deficit_reduction_ha",
    ]:
        df[col] = num(df, col)
    if "car_valid" in df.columns:
        df["car_valid"] = df["car_valid"].fillna(False).astype(bool)
    else:
        df["car_valid"] = True
    df["mun_geocodigo_norm"] = mun_code(df)
    return df


def load_priority_property_geometries() -> gpd.GeoDataFrame:
    fc = pd.read_parquet(
        today_path("forest_code_mt_priority_consolidated", "parquet"),
        columns=["priority_key", "input_file_type", "prop_id_unique", "car_join", "NUMEROESTA"],
    )
    geometry_sources = {
        "simcar_validado": RAW_DIR / "simcar_validado" / "CAR_ATP.parquet",
        "simcar_digital": PROCESSED_DIR / "simcar_digital" / "simcar_digital_march2026_geo_master.parquet",
        "simcar_proxy": PROCESSED_DIR / "simcar_proxy" / "simcar_p_march2026_geo_master.parquet",
    }
    frames = []
    for source, path in geometry_sources.items():
        sub = fc[fc["input_file_type"].eq(source)].copy()
        if sub.empty or not path.exists():
            continue
        cols = fcm_gta.fc_priority.available_columns(path)
        join_cols = [c for c in ["prop_id_unique", "car_join", "NUMEROESTA", "geometry"] if c in cols]
        if "geometry" not in join_cols:
            continue
        geom = gpd.read_parquet(path, columns=join_cols)
        merged = None
        if "prop_id_unique" in geom.columns and sub["prop_id_unique"].notna().any():
            merged = sub.merge(geom.drop_duplicates("prop_id_unique"), on="prop_id_unique", how="left")
        if (merged is None or merged["geometry"].isna().all()) and "car_join" in geom.columns and sub["car_join"].notna().any():
            merged = sub.merge(geom.drop_duplicates("car_join"), on="car_join", how="left")
        if merged is None or merged["geometry"].isna().all():
            if "prop_id_unique" in geom.columns:
                geom = geom.copy()
                geom["priority_key"] = geom["prop_id_unique"].astype(str).str.upper().str.extract(CAR_CODE_RE.pattern, expand=False)
                merged = sub.merge(geom.drop_duplicates("priority_key"), on="priority_key", how="left")
            else:
                continue
        if "NUMEROESTA" in geom.columns and merged["geometry"].isna().any():
            fallback = merged.loc[merged["geometry"].isna(), ["priority_key", "input_file_type"]]
            fallback = fallback.merge(sub[["priority_key", "NUMEROESTA"]].drop_duplicates("priority_key"), on="priority_key", how="left")
            geom_by_num = geom.drop(columns=["priority_key"], errors="ignore").dropna(subset=["NUMEROESTA"]).drop_duplicates("NUMEROESTA")
            fallback = fallback.merge(geom_by_num, on="NUMEROESTA", how="left")
            if not fallback.empty:
                merged = pd.concat([merged.loc[merged["geometry"].notna(), ["priority_key", "input_file_type", "geometry"]], fallback[["priority_key", "input_file_type", "geometry"]]], ignore_index=True)
        frames.append(gpd.GeoDataFrame(merged[["priority_key", "input_file_type", "geometry"]], geometry="geometry", crs=geom.crs))
    if not frames:
        return gpd.GeoDataFrame(columns=["priority_key", "input_file_type", "geometry"], geometry="geometry", crs=DISPLAY_CRS)
    gdf = pd.concat(frames, ignore_index=True, sort=False)
    gdf = gpd.GeoDataFrame(gdf, geometry="geometry", crs=frames[0].crs)
    return gdf[gdf.geometry.notna() & ~gdf.geometry.is_empty].drop_duplicates("priority_key").copy()


def secondary_vegetation_by_property(force: bool = False) -> pd.DataFrame:
    if SECONDARY_BY_PROPERTY.exists() and not force:
        return pd.read_parquet(SECONDARY_BY_PROPERTY)
    if not SECONDARY_FOREST.exists():
        return pd.DataFrame(columns=["priority_key", "secondary_vegetation_ha"])
    props = load_priority_property_geometries()
    secondary = gpd.read_file(SECONDARY_FOREST)
    if "layer" in secondary.columns:
        layer_mask = secondary["layer"].astype(str).eq("input_secondary_forest_dissolved")
        if layer_mask.any():
            secondary = secondary[layer_mask].copy()
    if secondary.empty or props.empty:
        return pd.DataFrame(columns=["priority_key", "secondary_vegetation_ha"])
    if secondary.crs != props.crs:
        secondary = secondary.to_crs(props.crs)
    props = props[["priority_key", "geometry"]].copy().reset_index(drop=True)
    secondary = secondary[secondary.geometry.notna() & ~secondary.geometry.is_empty][["geometry"]].copy().reset_index(drop=True)
    if secondary.empty:
        out = pd.DataFrame(columns=["priority_key", "secondary_vegetation_ha"])
    else:
        pair_idx = props.sindex.query(secondary.geometry, predicate="intersects")
        if pair_idx.size == 0:
            out = pd.DataFrame(columns=["priority_key", "secondary_vegetation_ha"])
        else:
            sec_idx = pair_idx[0]
            prop_idx = pair_idx[1]
            totals = np.zeros(len(props), dtype="float64")
            chunk_size = 100_000
            prop_geoms = props.geometry.array
            sec_geoms = secondary.geometry.array
            for start in range(0, len(prop_idx), chunk_size):
                end = start + chunk_size
                pidx = prop_idx[start:end]
                sidx = sec_idx[start:end]
                intersections = shapely.intersection(prop_geoms.take(pidx), sec_geoms.take(sidx))
                areas = shapely.area(intersections) / 10_000.0
                np.add.at(totals, pidx, np.nan_to_num(areas, nan=0.0, posinf=0.0, neginf=0.0))
            out = props.loc[totals > 0, ["priority_key"]].copy()
            out["secondary_vegetation_ha"] = totals[totals > 0]
            out = out.groupby("priority_key", as_index=False)["secondary_vegetation_ha"].sum().sort_values("priority_key")
    out.to_parquet(SECONDARY_BY_PROPERTY, index=False)
    return out


def table_cons2000_compliance(df: pd.DataFrame, group_cols: list[str] | None = None) -> pd.DataFrame:
    work = compliance.add_cons2000_scenarios(df)
    group_cols = group_cols or []
    if not group_cols:
        work = work.assign(Group="All properties")
        group_cols = ["Group"]
    rows = []
    for scenario, suffix, flag in [
        ("With 2000 rule", "", "compliant_with_cons2000"),
        ("Without 2000 rule", "_without_cons2000", "compliant_without_cons2000"),
    ]:
        rl_adj = "rl_adj_deficit_ha" if suffix == "" else "rl_adj_deficit_without_cons2000_ha"
        rl_restore = "rl_restore_ha" if suffix == "" else "rl_restore_without_cons2000_ha"
        rl_comp = "rl_compensate_ha" if suffix == "" else "rl_compensate_without_cons2000_ha"
        app_restore = "app_restore_ha" if suffix == "" else "app_restore_without_cons2000_ha"
        combined = "combined_with_cons2000_ha" if suffix == "" else "combined_without_cons2000_ha"
        sub = (
            work.groupby(group_cols, observed=False, dropna=False)
            .agg(
                Properties=("priority_key", "size"),
                **{
                    "Compliant (n)": (flag, "sum"),
                    "Any liability (n)": (flag, lambda s: int((~s.astype(bool)).sum())),
                    "LR Adj. Deficit (ha)": (rl_adj, "sum"),
                    "APP Restore (ha)": (app_restore, "sum"),
                    "Combined Liability (ha)": (combined, "sum"),
                    "To Restore (ha)": (rl_restore, "sum"),
                    "To Compensate (ha)": (rl_comp, "sum"),
                },
            )
            .reset_index()
        )
        sub["Scenario"] = scenario
        rows.append(sub)
    out = pd.concat(rows, ignore_index=True)
    props = pd.to_numeric(out["Properties"], errors="coerce").replace(0, np.nan)
    out["% Compliant"] = (pd.to_numeric(out["Compliant (n)"], errors="coerce").fillna(0) / props * 100).fillna(0)
    if len(group_cols) > 0:
        out = out.sort_values([*group_cols, "Scenario"], kind="stable")
    ordered = [*group_cols, "Scenario", "Properties", "Compliant (n)", "% Compliant", "Any liability (n)", "LR Adj. Deficit (ha)", "APP Restore (ha)", "Combined Liability (ha)", "To Restore (ha)", "To Compensate (ha)"]
    rename = {
        "size_class": "Size Class",
        "supplier_type": "Supplier Type",
        "Group": "Group",
        "Properties": "N",
        "Compliant (n)": "Compliant n",
        "Any liability (n)": "Liability n",
        "LR Adj. Deficit (ha)": "LR Def. (ha)",
        "APP Restore (ha)": "APP Liab. (ha)",
        "Combined Liability (ha)": "Total Liab. (ha)",
        "To Restore (ha)": "Restore (ha)",
        "To Compensate (ha)": "Comp. (ha)",
    }
    return out[ordered].rename(columns=rename)


def table_cons2000_delta(df: pd.DataFrame, group_cols: list[str] | None = None) -> pd.DataFrame:
    work = compliance.add_cons2000_scenarios(df)
    group_cols = group_cols or []
    if not group_cols:
        work = work.assign(Group="All properties")
        group_cols = ["Group"]
    out = (
        work.groupby(group_cols, observed=False, dropna=False)
        .agg(
            Properties=("priority_key", "size"),
            **{
                "Baseline LR liability (ha)": ("rl_adj_deficit_ha", "sum"),
                "Without 2000 LR liability (ha)": ("rl_adj_deficit_without_cons2000_ha", "sum"),
                "LR delta (ha)": ("rl_cons2000_delta_ha", "sum"),
                "Baseline APP liability (ha)": ("app_restore_ha", "sum"),
                "Without 2000 APP liability (ha)": ("app_restore_without_cons2000_ha", "sum"),
                "APP delta (ha)": ("app_cons2000_delta_ha", "sum"),
                "Baseline total liability (ha)": ("combined_with_cons2000_ha", "sum"),
                "Without 2000 total liability (ha)": ("combined_without_cons2000_ha", "sum"),
                "Total delta (ha)": ("combined_cons2000_delta_ha", "sum"),
            },
        )
        .reset_index()
    )
    rename = {
        "size_class": "Size Class",
        "supplier_type": "Supplier Type",
        "Group": "Group",
    }
    return out.rename(columns=rename)


def table_secondary_vegetation_impact(df: pd.DataFrame, group_cols: list[str] | None = None) -> pd.DataFrame:
    group_cols = group_cols or []
    work = df.copy()
    if "input_file_type" in work.columns:
        work["input_file_type"] = pd.Categorical(work["input_file_type"].astype(str), INPUT_ORDER, ordered=True)
    if not group_cols:
        work = work.assign(Group="All properties")
        group_cols = ["Group"]
    grouped = (
        work.groupby(group_cols, observed=False, dropna=False)
        .agg(
            Properties=("priority_key", "size"),
            Secondary_vegetation_ha=("secondary_vegetation_ha", "sum"),
            Properties_with_secondary=("secondary_vegetation_ha", lambda s: int((pd.to_numeric(s, errors="coerce").fillna(0) > 0).sum())),
            Baseline_compliant=("compliant_baseline", "sum"),
            With_secondary_compliant=("compliant_with_secondary", "sum"),
            Changed_to_compliant=("secondary_changes_to_compliant", "sum"),
            Baseline_total_liability_ha=("calc_deficit_total_ha", "sum"),
            With_secondary_total_liability_ha=("calc_deficit_total_with_secondary_ha", "sum"),
            Liability_reduction_ha=("secondary_deficit_reduction_ha", "sum"),
            Baseline_LR_deficit_ha=("rl_adj_deficit_ha", "sum"),
            With_secondary_LR_deficit_ha=("rl_adj_deficit_with_secondary_ha", "sum"),
            Baseline_restore_ha=("rl_restore_ha", "sum"),
            With_secondary_restore_ha=("rl_restore_with_secondary_ha", "sum"),
            Baseline_compensate_ha=("rl_compensate_ha", "sum"),
            With_secondary_compensate_ha=("rl_compensate_with_secondary_ha", "sum"),
        )
        .reset_index()
    )
    props = pd.to_numeric(grouped["Properties"], errors="coerce").replace(0, np.nan)
    grouped["% Properties with secondary veg."] = (grouped["Properties_with_secondary"] / props * 100).fillna(0)
    grouped["Baseline % compliant"] = (grouped["Baseline_compliant"] / props * 100).fillna(0)
    grouped["With secondary % compliant"] = (grouped["With_secondary_compliant"] / props * 100).fillna(0)
    grouped["Compliance gain (p.p.)"] = grouped["With secondary % compliant"] - grouped["Baseline % compliant"]
    rename = {
        "size_class": "Size Class",
        "supplier_type": "Supplier Type",
        "input_file_type": "Input File Type",
        "Group": "Group",
        "Properties": "N",
        "Secondary_vegetation_ha": "Secondary veg. (ha)",
        "Properties_with_secondary": "Secondary veg. n",
        "Baseline_compliant": "Baseline compliant n",
        "With_secondary_compliant": "With secondary compliant n",
        "Changed_to_compliant": "Newly compliant n",
        "Baseline_total_liability_ha": "Baseline total liab. (ha)",
        "With_secondary_total_liability_ha": "With secondary total liab. (ha)",
        "Liability_reduction_ha": "Liab. reduction (ha)",
        "Baseline_LR_deficit_ha": "Baseline LR def. (ha)",
        "With_secondary_LR_deficit_ha": "With secondary LR def. (ha)",
        "Baseline_restore_ha": "Baseline restore (ha)",
        "With_secondary_restore_ha": "With secondary restore (ha)",
        "Baseline_compensate_ha": "Baseline comp. (ha)",
        "With_secondary_compensate_ha": "With secondary comp. (ha)",
    }
    ordered = [
        *group_cols,
        "Properties",
        "Secondary_vegetation_ha",
        "Properties_with_secondary",
        "% Properties with secondary veg.",
        "Baseline_compliant",
        "Baseline % compliant",
        "With_secondary_compliant",
        "With secondary % compliant",
        "Changed_to_compliant",
        "Compliance gain (p.p.)",
        "Baseline_total_liability_ha",
        "With_secondary_total_liability_ha",
        "Liability_reduction_ha",
        "Baseline_LR_deficit_ha",
        "With_secondary_LR_deficit_ha",
        "Baseline_restore_ha",
        "With_secondary_restore_ha",
        "Baseline_compensate_ha",
        "With_secondary_compensate_ha",
    ]
    result = grouped[ordered].rename(columns=rename)
    if "Input File Type" in result.columns:
        result["Input File Type"] = pd.Categorical(result["Input File Type"].astype(str), INPUT_ORDER, ordered=True)
        result = result.sort_values([c for c in ["Input File Type", "Size Class"] if c in result.columns], kind="stable")
        result["Input File Type"] = display_input_labels(result["Input File Type"])
    return result


def table_status(df: pd.DataFrame) -> pd.DataFrame:
    total_n = len(df)
    total_area = df["area_ha_car"].sum()
    rows = []
    for label, sub in [("Active", df[df["car_valid"]]), ("Inactive (Canceled / Rejected)", df[~df["car_valid"]]), ("Total", df)]:
        rows.append({
            "Registration Status": label,
            "Number of Properties": len(sub),
            "% of Properties": pct(len(sub), total_n),
            "Area (Mha)": sub["area_ha_car"].sum() / 1e6,
            "% of Area": pct(sub["area_ha_car"].sum(), total_area),
        })
    return pd.DataFrame(rows)


def table_status_detail(df: pd.DataFrame) -> pd.DataFrame:
    total_n = len(df)
    total_area = df["area_ha_car"].sum()
    return (
        df.groupby("SITUACAO", dropna=False)
        .agg(N=("priority_key", "size"), Area_ha=("area_ha_car", "sum"))
        .reset_index()
        .assign(**{"% N": lambda d: pct(d["N"], total_n), "% Area": lambda d: pct(d["Area_ha"], total_area)})
        .sort_values("N", ascending=False)
    )


def table_size(df: pd.DataFrame) -> pd.DataFrame:
    total_n = len(df)
    total_area = df["area_ha_car"].sum()
    out = (
        df.groupby("size_class", observed=False)
        .agg(N=("priority_key", "size"), **{"Area (ha)": ("area_ha_car", "sum")})
        .reset_index()
        .assign(**{"% N": lambda d: pct(d["N"], total_n), "% Area": lambda d: pct(d["Area (ha)"], total_area)})
    )
    total = pd.DataFrame([{"size_class": "Total", "N": total_n, "Area (ha)": total_area, "% N": 100.0, "% Area": 100.0}])
    return pd.concat([out, total], ignore_index=True).rename(columns={"size_class": "Size Class"})


def table_status_by_size(df: pd.DataFrame) -> pd.DataFrame:
    return (
        df.groupby(["size_class", "car_valid"], observed=False)
        .agg(N=("priority_key", "size"), Area_ha=("area_ha_car", "sum"))
        .reset_index()
        .assign(Status=lambda d: np.where(d["car_valid"], "Active", "Inactive"))
        .drop(columns="car_valid")
        .rename(columns={"size_class": "Size Class"})
    )


def table_special_areas(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for label, cols in SPECIAL_AREA_SPECS.items():
        vals = sum((num(df, c) for c in cols), start=pd.Series(0.0, index=df.index))
        sub = df[vals > 0]
        rows.append({"Special Area": label, "N": len(sub), "Area overlap (ha)": vals.sum(), "Property area (ha)": sub["area_ha_car"].sum()})
    return pd.DataFrame(rows)


def table_lr(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    return (
        df.groupby(group_cols, observed=False, dropna=False)
        .agg(
            **{
                "Required (ha)": ("rl_req_total_ha", "sum"),
                "Gross Deficit": ("rl_gross_deficit_ha", "sum"),
                "Adj. Deficit": ("rl_adj_deficit_ha", "sum"),
                "To Restore": ("rl_restore_ha", "sum"),
                "To Compensate": ("rl_compensate_ha", "sum"),
                "Surplus (ha)": ("rl_surplus_total_ha", "sum"),
            }
        )
        .reset_index()
    )


def table_app(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    work = df.copy()
    work["app_net_deficit_ha"] = work["app_restore_ha"]
    work["has_app_net_deficit"] = work["app_net_deficit_ha"] > 0
    return (
        work.groupby(group_cols, observed=False, dropna=False)
        .agg(
            **{
                "APP Required (ha)": ("app_req_ha", "sum"),
                "APP Preserved (ha)": ("app_preserved_ha", "sum"),
                "APP Gross Deficit (ha)": ("app_gross_deficit_ha", "sum"),
                "APP Net Deficit (ha)": ("app_net_deficit_ha", "sum"),
                "Mean Net deficit (ha)": ("app_net_deficit_ha", lambda s: s[s > 0].mean() if (s > 0).any() else 0),
                "Properties Net Deficit (n)": ("has_app_net_deficit", "sum"),
            }
        )
        .reset_index()
    )


def table_combined(df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    work = df.copy()
    work["combined_gross"] = work["rl_gross_deficit_ha"] + work["app_gross_deficit_ha"]
    work["combined_adj"] = work["rl_adj_deficit_ha"] + work["app_restore_ha"]
    work["combined_restore"] = work["rl_restore_ha"] + work["app_restore_ha"]
    return (
        work.groupby(group_cols, observed=False, dropna=False)
        .agg(
            **{
                "Gross Deficit (ha)": ("combined_gross", "sum"),
                "Adj. Deficit (ha)": ("combined_adj", "sum"),
                "To Restore (ha)": ("combined_restore", "sum"),
                "To Compensate (ha)": ("rl_compensate_ha", "sum"),
            }
        )
        .reset_index()
    )


def table_mean_deficit(df: pd.DataFrame) -> pd.DataFrame:
    work = df.copy()
    work["total_mean_deficit"] = work["rl_adj_deficit_ha"] + work["app_restore_ha"]
    return (
        work.groupby("size_class", observed=False)
        .agg(**{"Total Mean Deficit (ha)": ("total_mean_deficit", lambda s: s[s > 0].mean() if (s > 0).any() else 0)})
        .reset_index()
        .rename(columns={"size_class": "Size Class"})
    )


def table_top_municipalities(df: pd.DataFrame, n: int = 20) -> pd.DataFrame:
    mun_names = load_municipality_names()
    out = (
        df.groupby("mun_geocodigo_norm", dropna=False)
        .agg(
            Properties=("priority_key", "size"),
            **{
                "LR Surplus": ("rl_surplus_total_ha", "sum"),
                "LR Deficit": ("rl_adj_deficit_ha", "sum"),
                "APP Deficit": ("app_restore_ha", "sum"),
            },
        )
        .reset_index()
    )
    out["Total Deficit"] = out["LR Deficit"] + out["APP Deficit"]
    out = out.merge(mun_names, on="mun_geocodigo_norm", how="left")
    out["Municipality"] = out["municipality_name"].fillna(out["mun_geocodigo_norm"])
    return out.sort_values("Total Deficit", ascending=False).head(n)[["Municipality", "Properties", "LR Surplus", "LR Deficit", "APP Deficit", "Total Deficit"]]


def load_municipality_names() -> pd.DataFrame:
    if not MUN_GEOM.exists():
        return pd.DataFrame(columns=["mun_geocodigo_norm", "municipality_name"])
    gdf = gpd.read_parquet(MUN_GEOM, columns=["nome", "geocodigo", "geometry"])
    return (
        pd.DataFrame({"mun_geocodigo_norm": gdf["geocodigo"].astype("string").str.zfill(7), "municipality_name": gdf["nome"].astype(str)})
        .drop_duplicates("mun_geocodigo_norm")
    )


def table_cattle_characterization(joined: pd.DataFrame) -> pd.DataFrame:
    total_n = len(joined)
    total_area = joined["area_ha_car"].sum()
    out = (
        joined.groupby(["supplier_type", "size_class"], observed=False, dropna=False)
        .agg(N=("priority_key", "size"), **{"Area (ha)": ("area_ha_car", "sum"), "Cattle head": ("total_cattle_moved", "sum")})
        .reset_index()
        .assign(**{"% N": lambda d: pct(d["N"], total_n), "% Area": lambda d: pct(d["Area (ha)"], total_area)})
        .rename(columns={"supplier_type": "Category", "size_class": "Size Class"})
    )
    out["Category"] = pd.Categorical(out["Category"], SUPPLIER_ORDER, ordered=True)
    return out.sort_values(["Category", "Size Class"])


def table_cattle_supplier_subgroups(joined: pd.DataFrame) -> pd.DataFrame:
    work = joined.copy()
    if "supplier_macro" not in work.columns:
        work["supplier_macro"] = work["supplier_type"]
    if "supplier_subgroup" not in work.columns:
        work["supplier_subgroup"] = np.select(
            [
                (num(work, "slaughter_or_export_cattle") > 0) & (num(work, "total_t1_cattle") > 0),
                num(work, "slaughter_or_export_cattle") > 0,
                (num(work, "slaughter_or_export_cattle") == 0) & (num(work, "total_t1_cattle") > 0),
            ],
            ["Direct also Tier 1", "Direct pure", "Tier 1 pure"],
            default="Tier 2+",
        )
    work["direct_gt50_slaughter"] = (
        work["supplier_type"].astype(str).eq("Direct supplier") &
        (num(work, "total_cattle_slaughter") / num(work, "total_cattle_moved").replace(0, np.nan) > 0.5)
    )
    out = (
        work.groupby(["supplier_macro", "supplier_subgroup"], observed=False, dropna=False)
        .agg(
            Properties=("priority_key", "size"),
            **{
                "Cattle Head": ("total_cattle_moved", "sum"),
                "Slaughter Head": ("total_cattle_slaughter", "sum"),
                "Slaughter/Export Head": ("slaughter_or_export_cattle", "sum"),
                "Direct >50% Slh (n)": ("direct_gt50_slaughter", "sum"),
            },
        )
        .reset_index()
        .rename(columns={"supplier_macro": "Macro Class", "supplier_subgroup": "Subgroup"})
    )
    out["Macro Class"] = pd.Categorical(out["Macro Class"], SUPPLIER_ORDER, ordered=True)
    subgroup_order = ["Direct pure", "Direct also Tier 1", "Tier 1 pure", "Tier 2+"]
    out["Subgroup"] = pd.Categorical(out["Subgroup"], subgroup_order, ordered=True)
    out = out.sort_values(["Macro Class", "Subgroup"]).reset_index(drop=True)
    total = pd.DataFrame([{
        "Macro Class": "Total",
        "Subgroup": "Total",
        "Properties": out["Properties"].sum(),
        "Cattle Head": out["Cattle Head"].sum(),
        "Slaughter Head": out["Slaughter Head"].sum(),
        "Slaughter/Export Head": out["Slaughter/Export Head"].sum(),
        "Direct >50% Slh (n)": out["Direct >50% Slh (n)"].sum(),
    }])
    return pd.concat([out, total], ignore_index=True)


def table_net_liabilities(joined: pd.DataFrame) -> pd.DataFrame:
    return (
        joined.groupby("supplier_type", observed=False, dropna=False)
        .agg(**{"To Restore (ha)": ("rl_restore_ha", "sum"), "APP Restore (ha)": ("app_restore_ha", "sum"), "To Compensate (ha)": ("rl_compensate_ha", "sum")})
        .reset_index()
        .rename(columns={"supplier_type": "Category"})
    )


def add_total_row(df: pd.DataFrame, label_col: str, label: str = "Total") -> pd.DataFrame:
    numeric_cols = df.select_dtypes(include=[np.number]).columns
    row = {label_col: label}
    for col in numeric_cols:
        row[col] = df[col].sum()
    return pd.concat([df, pd.DataFrame([row])], ignore_index=True)


def table_database_sources() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "Input data": "SIMCAR validated",
            "Priority": 1,
            "Analytical use": "Validated CAR property layer",
            "Short path": "simcar_validado",
            "Processing rule": "Priority 1; removes duplicates from digital/proxy.",
        },
        {
            "Input data": "SIMCAR digital",
            "Priority": 2,
            "Analytical use": "Digital CAR property layer",
            "Short path": "simcar_digital",
            "Processing rule": "Priority 2; used where no validated CAR exists.",
        },
        {
            "Input data": "SIMCAR proxy",
            "Priority": 3,
            "Analytical use": "Proxy CAR property layer",
            "Short path": "fc_summary_mt",
            "Processing rule": "Priority 3; remaining farms only.",
        },
        {
            "Input data": "GTA Mato Grosso",
            "Priority": "NA",
            "Analytical use": "Cattle supplier classification",
            "Short path": "gta/erich_car_gta",
            "Processing rule": "Direct if slaughter/export > 0; direct overrides Tier 1.",
        },
        {
            "Input data": "PRODES 2000",
            "Priority": "NA",
            "Analytical use": "Consolidated area 2000",
            "Short path": "cons_area_2000",
            "Processing rule": "MMU 6.25 hectares; fragments < 6.25 hectares removed.",
        },
        {
            "Input data": "Secondary vegetation",
            "Priority": "NA",
            "Analytical use": "Scenario impact on LR compliance and total liabilities",
            "Short path": "input_secondary_forest_dissolved",
            "Processing rule": "Intersected spatially with prioritized validado, digital and proxy CAR geometries; reported as baseline vs including secondary vegetation.",
        },
    ])


def with_category(df: pd.DataFrame, category: str, label_col: str = "Size Class") -> pd.DataFrame:
    out = df.copy()
    out.insert(0, "Category", category)
    total_mask = out[label_col].astype(str).str.lower().eq("total")
    out.loc[total_mask, "Category"] = "Total"
    return out


def add_share_columns(df: pd.DataFrame, metrics: list[tuple[str, str]]) -> pd.DataFrame:
    out = df.copy()
    for value_col, pct_col in metrics:
        denom = pd.to_numeric(out.loc[out["Category"].ne("Total"), value_col], errors="coerce").fillna(0).sum()
        out[pct_col] = pct(pd.to_numeric(out[value_col], errors="coerce").fillna(0), denom) if denom else 0.0
    ordered = ["Category", "Size Class"]
    for value_col, pct_col in metrics:
        ordered.extend([value_col, pct_col])
    return out[ordered]


def table_size_category(df: pd.DataFrame, category: str) -> pd.DataFrame:
    return with_category(table_size(df), category)


def table_lr_category(df: pd.DataFrame, category: str) -> pd.DataFrame:
    base = table_lr(df, ["size_class"]).rename(columns={
        "size_class": "Size Class",
        "Required (ha)": "LR Req",
        "Gross Deficit": "Grs Def",
        "Adj. Deficit": "Adj Def",
        "To Restore": "Rest.",
        "To Compensate": "Comp.",
        "Surplus (ha)": "Surplus",
    })
    base = add_total_row(base, "Size Class")
    return add_share_columns(with_category(base, category), [
        ("LR Req", "% Req"),
        ("Grs Def", "% Grs"),
        ("Adj Def", "% Adj"),
        ("Rest.", "% Rest"),
        ("Comp.", "% Comp"),
        ("Surplus", "% Sur"),
    ])


def table_app_category(df: pd.DataFrame, category: str) -> pd.DataFrame:
    base = table_app(df, ["size_class"]).rename(columns={
        "size_class": "Size Class",
        "APP Required (ha)": "APP Req",
        "APP Preserved (ha)": "APP Pres.",
        "APP Gross Deficit (ha)": "Grs Def",
        "APP Net Deficit (ha)": "Net Def",
    })
    keep = ["Size Class", "APP Req", "APP Pres.", "Grs Def", "Net Def"]
    base = add_total_row(base[keep], "Size Class")
    return add_share_columns(with_category(base, category), [
        ("APP Req", "% Req"),
        ("APP Pres.", "% Pres"),
        ("Grs Def", "% Grs"),
        ("Net Def", "% Net"),
    ])


def table_combined_category(df: pd.DataFrame, category: str) -> pd.DataFrame:
    base = table_combined(df, ["size_class"]).rename(columns={
        "size_class": "Size Class",
        "Gross Deficit (ha)": "Grs Def",
        "Adj. Deficit (ha)": "Adj Def",
        "To Restore (ha)": "Rest.",
        "To Compensate (ha)": "Comp.",
    })
    base = add_total_row(base, "Size Class")
    return add_share_columns(with_category(base, category), [
        ("Grs Def", "% Grs"),
        ("Adj Def", "% Adj"),
        ("Rest.", "% Rest"),
        ("Comp.", "% Comp"),
    ])


def special_area_subset(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    vals = sum((num(df, c) for c in cols), start=pd.Series(0.0, index=df.index))
    return df[vals > 0].copy()


def concat_category_tables(parts: list[pd.DataFrame], subtotal_label: str = "Subtotal") -> pd.DataFrame:
    rows = []
    for part in parts:
        if part.empty:
            continue
        category = str(part["Category"].iloc[0])
        body = part[part["Category"].ne("Total")].copy()
        subtotal = part[part["Category"].eq("Total")].copy()
        if not subtotal.empty:
            subtotal["Category"] = subtotal_label
        body["Category"] = category
        rows.extend([body, subtotal])
    out = pd.concat(rows, ignore_index=True) if rows else pd.DataFrame()
    if out.empty:
        return out
    numeric_cols = out.select_dtypes(include=[np.number]).columns
    total = {col: out.loc[out["Category"].eq(subtotal_label), col].sum() for col in numeric_cols}
    total["Category"] = "Total"
    total["Size Class"] = "Total"
    return recompute_percent_columns(pd.concat([out, pd.DataFrame([total])], ignore_index=True))


def recompute_percent_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    total_rows = out["Category"].astype(str).str.lower().eq("total") if "Category" in out.columns else pd.Series(False, index=out.index)
    for idx, col in enumerate(out.columns):
        if not str(col).startswith("%") or idx == 0:
            continue
        value_col = out.columns[idx - 1]
        values = pd.to_numeric(out[value_col], errors="coerce").fillna(0)
        denom = values[total_rows].iloc[-1] if total_rows.any() else values.sum()
        out[col] = values / denom * 100 if denom else 0.0
    return out


def table_special_by_size(df: pd.DataFrame) -> pd.DataFrame:
    parts = [table_size_category(special_area_subset(df, cols), label) for label, cols in SPECIAL_AREA_SPECS.items()]
    return concat_category_tables(parts)


def table_special_metric(df: pd.DataFrame, fn) -> pd.DataFrame:
    parts = [fn(special_area_subset(df, cols), label) for label, cols in SPECIAL_AREA_SPECS.items()]
    return concat_category_tables(parts)


def table_cattle_size(joined: pd.DataFrame) -> pd.DataFrame:
    labels = {
        "Direct supplier": "Direct Supplier",
        "Indirect supplier - Tier 1": "Indirect T1",
        "Indirect supplier - Tier 2+": "Indirect T2+",
    }
    parts = []
    for raw, label in labels.items():
        parts.append(table_size_category(joined[joined["supplier_type"].astype(str).eq(raw)], label))
    return concat_category_tables(parts)


def table_cattle_metric(joined: pd.DataFrame, fn) -> pd.DataFrame:
    labels = {
        "Direct supplier": "Direct Supp.",
        "Indirect supplier - Tier 1": "Ind. T1",
        "Indirect supplier - Tier 2+": "Ind. T2+",
    }
    parts = []
    for raw, label in labels.items():
        parts.append(fn(joined[joined["supplier_type"].astype(str).eq(raw)], label))
    return concat_category_tables(parts)


def save_stacked_bar(data: pd.DataFrame, x: str, cols: list[str], title: str, path: Path) -> Path:
    data = data.loc[data[cols].fillna(0).sum(axis=1).gt(0)].copy()
    fig, ax = plt.subplots(figsize=CHART_SIZE_IN)
    fig.patch.set_facecolor("white")
    bottom = np.zeros(len(data))
    colors = ["#D8743C", "#6B7378", "#AEB5B9", "#D8DCDE"]
    labels = data[x].astype("object").where(pd.notna(data[x]), "Not Classified").astype(str).tolist()
    for col, color in zip(cols, colors):
        ax.bar(labels, data[col], bottom=bottom, label=col, color=color)
        bottom += data[col].to_numpy()
    polish_chart(fig, ax, "Area (ha)")
    add_chart_title(ax, title)
    add_bar_labels(ax, stacked_totals=bottom)
    place_legend_right(ax)
    horizontal_xticks(ax, labels, width=14)
    square_chart_layout(fig)
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def save_grouped_bar(data: pd.DataFrame, x: str, y: str, hue: str, title: str, path: Path) -> Path:
    data = data.copy()
    data[x] = data[x].astype("object").where(pd.notna(data[x]), "Not Classified").astype(str)
    data[hue] = data[hue].astype("object").where(pd.notna(data[hue]), "Other").astype(str)
    piv = data.pivot_table(index=x, columns=hue, values=y, aggfunc="sum", fill_value=0)
    piv = piv.loc[piv.sum(axis=1).gt(0), piv.sum(axis=0).gt(0)]
    ax = piv.plot(kind="bar", figsize=CHART_SIZE_IN, color=["#D8743C", "#59666D", "#AEB5B9", "#D8DCDE"])
    ax.figure.patch.set_facecolor("white")
    polish_chart(ax.figure, ax, y)
    add_chart_title(ax, title)
    add_bar_labels(ax)
    place_legend_right(ax)
    horizontal_xticks(ax, piv.index.astype(str), width=14)
    square_chart_layout(ax.figure)
    ax.figure.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(ax.figure)
    return path


def save_pie_counts(data: pd.DataFrame, labels: pd.Series, title: str, path: Path) -> Path:
    fig, ax = plt.subplots(figsize=(10.0, 2.2))
    fig.patch.set_facecolor("white")
    label_text = [f"{label} — {compact_number(value)}" for label, value in zip(labels, data)]
    values = np.asarray(data, dtype=float)
    total = values.sum()
    left = 0.0
    for value, label, color in zip(values, labels, ["#7A858B", "#D8743C", "#AEB5B9", "#D8DCDE"]):
        share = value / total * 100 if total else 0
        ax.barh([0], [share], left=left, color=color, height=0.42)
        ax.text(left + share / 2, 0, f"{label}\n{value:,.0f} ({share:.1f}%)",
                ha="center", va="center", fontsize=FIG_FONT_SIZE, color="white", fontweight="bold")
        left += share
    ax.set_xlim(0, 100)
    ax.set_ylim(-0.23, 0.23)
    ax.set_axis_off()
    add_chart_title(ax, title)
    fig.subplots_adjust(left=.06, right=.98, bottom=.12, top=.96)
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def make_pdf_style_maps(df: pd.DataFrame) -> list[Path]:
    if not MUN_GEOM.exists():
        return []
    apply_map_style()
    colors = map_colors()
    theme = map_theme()
    mun = gpd.read_parquet(MUN_GEOM)
    mun = mun[mun["geocodigo"].astype(str).str.startswith("51")].copy()
    mt = gpd.read_parquet(UF_GEOM)
    mt = mt[mt["sigla"].astype(str).eq("MT")].copy()
    if mun.crs is not None:
        mun = mun.to_crs(DISPLAY_CRS)
    if mt.crs is not None:
        mt = mt.to_crs(DISPLAY_CRS)
    work = df.copy()
    work["any_deficit"] = (work["rl_adj_deficit_ha"] + work["app_restore_ha"]) > 0
    mun_tbl = (
        work.groupby("mun_geocodigo_norm", dropna=False)
        .agg(
            n_properties=("priority_key", "size"),
            mean_total_deficit_ha=("calc_deficit_total_ha", "mean"),
            pct_any_deficit=("any_deficit", lambda s: s.mean() * 100),
            total_deficit_ha=("calc_deficit_total_ha", "sum"),
            restoration_ha=("rl_restore_ha", "sum"),
        )
        .reset_index()
    )
    gdf = mun.merge(mun_tbl, left_on="geocodigo", right_on="mun_geocodigo_norm", how="left")
    maps = [
        ("mean_total_deficit_ha", "Mean total deficit per property", "YlOrRd", "hectares"),
        ("pct_any_deficit", "Properties with any deficit", "Purples", "%"),
        ("total_deficit_ha", "Aggregate total deficit", "OrRd", "hectares"),
        ("restoration_ha", "Restoration liability", "YlGn", "hectares"),
    ]
    paths = []
    fig = plt.figure(figsize=MAP_SIZE_IN, facecolor=colors["paper"])
    positions = [
        (0.055, 0.525, 0.33, 0.275),
        (0.545, 0.525, 0.33, 0.275),
        (0.055, 0.165, 0.33, 0.275),
        (0.545, 0.165, 0.33, 0.275),
    ]
    for idx, ((col, title, cmap, unit), pos) in enumerate(zip(maps, positions)):
        ax = fig.add_axes(pos)
        ax.set_facecolor("white")
        ax.set_title(title, loc="left", fontsize=FIG_FONT_SIZE, fontweight="bold", color=colors["ink"], pad=5)
        gdf[col] = pd.to_numeric(gdf[col], errors="coerce").fillna(0)
        plot = gdf.plot(
            column=col,
            ax=ax,
            cmap=cmap,
            linewidth=0.12,
            edgecolor=colors["boundary_light"],
            legend=False,
            missing_kwds={"color": colors["context_fill"]},
        )
        if not mt.empty:
            mt.boundary.plot(ax=ax, color=colors["ink"], linewidth=0.55)
        set_map_extent(ax, mt if not mt.empty else gdf, pad_x=0.015, pad_y=0.025)
        if idx == 0 and draw_scale_bar is not None and theme is not None:
            draw_scale_bar(ax, DISPLAY_CRS, "300 km", theme, length_m=300_000)
        sm = plt.cm.ScalarMappable(
            cmap=cmap,
            norm=plt.Normalize(vmin=float(gdf[col].min()), vmax=float(gdf[col].max())),
        )
        sm._A = []
        cax = fig.add_axes((pos[0] + pos[2] + 0.012, pos[1] + 0.02, 0.015, pos[3] - 0.04))
        cax.set_facecolor("white")
        cbar = fig.colorbar(sm, cax=cax)
        cbar.formatter = FuncFormatter(compact_percent if unit == "%" else compact_number)
        cbar.update_ticks()
        cbar.ax.tick_params(labelsize=FIG_FONT_SIZE, colors=colors["ink"], width=0.35)
        cbar.outline.set_linewidth(0.35)
        cbar.set_label(unit, fontsize=FIG_FONT_SIZE, color=colors["muted"], labelpad=2)
    add_map_source(fig, f"Source: Forest Code model outputs; IBGE BC250 municipal boundaries. CRS: {CRS_LABEL}.")
    panel_path = FIG_DIR / "Figure_09_municipal_noncompliance_panel.png"
    fig.savefig(panel_path, dpi=FIG_DPI, facecolor=fig.get_facecolor())
    plt.close(fig)
    paths.append(panel_path)
    return paths


def save_study_area_map() -> Path | None:
    if not MUN_GEOM.exists() or not UF_GEOM.exists():
        return None
    apply_map_style()
    colors = map_colors()
    theme = map_theme()
    mun = gpd.read_parquet(MUN_GEOM)
    mun = mun[mun["geocodigo"].astype(str).str.startswith("51")].copy()
    uf = gpd.read_parquet(UF_GEOM)
    mt = uf[uf["sigla"].astype(str).eq("MT")].copy()
    if mun.crs is not None:
        mun = mun.to_crs(DISPLAY_CRS)
    if mt.crs is not None:
        mt = mt.to_crs(DISPLAY_CRS)
    slots = map_layout("map-plus-metric")
    fig = plt.figure(figsize=MAP_SIZE_IN, facecolor=colors["paper"])
    ax = fig.add_axes((0.05, 0.08, 0.90, 0.88))
    ax.set_facecolor("white")
    mun.plot(ax=ax, color=colors["context_fill"], edgecolor=colors["boundary_light"], linewidth=0.18)
    if not mt.empty:
        mt.boundary.plot(ax=ax, color=colors["primary_edge"], linewidth=1.05)
    set_map_extent(ax, mt if not mt.empty else mun)
    if draw_scale_bar is not None and theme is not None:
        draw_scale_bar(ax, DISPLAY_CRS, "300 km", theme, length_m=300_000)
    area_mha = float(mt.geometry.area.sum() / 10_000 / 1_000_000) if not mt.empty else float(mun.geometry.area.sum() / 10_000 / 1_000_000)
    if False and draw_legend is not None and LegendItem is not None and theme is not None:
        draw_legend(
            fig,
            (
                LegendItem("Mato Grosso municipalities", colors["context_fill"], colors["boundary_light"]),
                LegendItem("State boundary", colors["paper"], colors["primary_edge"]),
            ),
            (slots.rail.left - 0.01, slots.rail.top - 0.18),
            theme,
        )
    add_map_source(fig, f"Source: IBGE BC250 municipal and federation-unit boundaries. CRS: {CRS_LABEL}.")
    path = FIG_DIR / "Figure_03_study_area_mato_grosso.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def save_vegetation_cover_figure(df: pd.DataFrame) -> Path:
    forest = df["radam_forest_ha"].sum()
    cerrado = df["radam_cerrado_ha"].sum()
    data = pd.DataFrame({"Vegetation class": ["Forest", "Cerrado"], "Area (Mha)": [forest / 1e6, cerrado / 1e6]})
    fig, ax = plt.subplots(figsize=CHART_SIZE_IN)
    fig.patch.set_facecolor("white")
    ax.bar(data["Vegetation class"], data["Area (Mha)"], color=["#356B58", "#D39A3A"])
    polish_chart(fig, ax, "Area (million ha)")
    add_chart_title(ax, "Figure 10: Native vegetation basis\nForest and Cerrado area intersecting active properties")
    for container in ax.containers:
        ax.bar_label(container, labels=[f"{bar.get_height():,.1f} million hectares" for bar in container], padding=2, fontsize=FIG_FONT_SIZE, color="#333333")
    fig.subplots_adjust(left=.11, right=.98, bottom=.14, top=.96)
    path = FIG_DIR / "Figure_10_vegetation_cover.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def save_secondary_impact_figure(df: pd.DataFrame) -> Path:
    data = table_secondary_vegetation_impact(df, ["input_file_type"]).rename(columns={"Input File Type": "Input"})
    data["Input"] = data["Input"].map({
        "simcar_validado": "SIMCAR validated",
        "simcar_digital": "SIMCAR digital",
        "simcar_proxy": "SIMCAR proxy",
    }).fillna(data["Input"])
    fig, ax = plt.subplots(figsize=CHART_SIZE_IN)
    fig.patch.set_facecolor("white")
    x = np.arange(len(data))
    width = 0.36
    baseline = data["Baseline total liab. (ha)"].to_numpy()
    with_sec = data["With secondary total liab. (ha)"].to_numpy()
    ax.bar(x - width / 2, baseline, width, label="Baseline", color="#AEB5B9")
    ax.bar(x + width / 2, with_sec, width, label="Including secondary vegetation", color="#D8743C")
    ax.set_xticks(x)
    ax.set_xticklabels(wrap_axis_labels(data["Input"], width=14), rotation=0, ha="center", fontsize=FIG_FONT_SIZE)
    polish_chart(fig, ax, "Total liability (ha)")
    add_chart_title(ax, "Figure 11: Secondary vegetation sensitivity\nEstimated liability with and without secondary vegetation")
    add_bar_labels(ax)
    place_legend_right(ax)
    square_chart_layout(fig)
    path = FIG_DIR / "Figure_11_secondary_vegetation_impact.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def save_cons2000_overall_figure(df: pd.DataFrame) -> Path:
    data = table_cons2000_delta(df)
    row = data.iloc[0]
    plot = pd.DataFrame({
        "Scenario": ["With 2000 rule", "Without 2000 rule"],
        "Legal Reserve (ha)": [row["Baseline LR liability (ha)"], row["Without 2000 LR liability (ha)"]],
        "APP (ha)": [row["Baseline APP liability (ha)"], row["Without 2000 APP liability (ha)"]],
    })
    fig, ax = plt.subplots(figsize=CHART_SIZE_IN)
    fig.patch.set_facecolor("white")
    bottom = np.zeros(len(plot))
    colors = ["#59666D", "#D8743C"]
    for idx, col in enumerate(["Legal Reserve (ha)", "APP (ha)"]):
        vals = plot[col].to_numpy(dtype=float)
        ax.bar(plot["Scenario"], vals, bottom=bottom, label=col.replace(" (ha)", ""), color=colors[idx])
        bottom += vals
    polish_chart(fig, ax, "Total liability (ha)")
    add_chart_title(ax, "Figure 12: Effect of the 2000 rule\nTotal liability under the two scenarios")
    for i, total in enumerate(bottom):
        ax.text(i, total, compact_number(total), ha="center", va="bottom", fontsize=FIG_FONT_SIZE, color="#333333")
    place_legend_right(ax)
    horizontal_xticks(ax, plot["Scenario"], width=14)
    square_chart_layout(fig)
    path = FIG_DIR / "Figure_12_cons2000_total_scenario.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def save_cons2000_size_figure(df: pd.DataFrame) -> Path:
    data = table_cons2000_delta(df, ["size_class"]).rename(columns={"Size Class": "Size"})
    labels = data["Size"].astype(str).tolist()
    x = np.arange(len(data))
    width = 0.36
    baseline = data["Baseline total liability (ha)"].to_numpy(dtype=float)
    without = data["Without 2000 total liability (ha)"].to_numpy(dtype=float)
    fig, ax = plt.subplots(figsize=CHART_SIZE_IN)
    fig.patch.set_facecolor("white")
    ax.bar(x - width / 2, baseline, width, label="With 2000 rule", color="#AEB5B9")
    ax.bar(x + width / 2, without, width, label="Without 2000 rule", color="#D8743C")
    ax.set_xticks(x)
    ax.set_xticklabels(wrap_axis_labels(labels, width=14), rotation=0, ha="center", fontsize=FIG_FONT_SIZE)
    polish_chart(fig, ax, "Total liability (ha)")
    add_chart_title(ax, "Figure 13: Effect of the 2000 rule by property size\nTotal liability under the two scenarios")
    add_bar_labels(ax)
    place_legend_right(ax)
    square_chart_layout(fig)
    path = FIG_DIR / "Figure_13_cons2000_size_scenario.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def save_supplier_subgroups_figure(joined: pd.DataFrame) -> Path:
    data = table_cattle_supplier_subgroups(joined)
    data = data[data["Macro Class"].astype(str).ne("Total")].copy()
    data["Label"] = data["Subgroup"].astype(str)
    x = np.arange(len(data))
    width = 0.38
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=CHART_SIZE_IN, sharey=True)
    fig.patch.set_facecolor("white")
    y = np.arange(len(data))
    ax1.barh(y, data["Properties"].to_numpy(dtype=float), color="#59666D")
    ax2.barh(y, data["Cattle Head"].to_numpy(dtype=float), color="#D8743C")
    ax1.set_yticks(y, wrap_axis_labels(data["Label"], width=22), fontsize=FIG_FONT_SIZE)
    ax1.invert_yaxis()
    for ax, label in [(ax1, "Properties"), (ax2, "Cattle head")]:
        ax.xaxis.set_major_formatter(FuncFormatter(abbreviated_number))
        ax.set_xlabel(label, fontsize=FIG_FONT_SIZE)
        ax.tick_params(labelsize=FIG_FONT_SIZE)
        ax.grid(axis="x", color="#DCE4E8", linewidth=.7)
        ax.spines[["top", "right", "left"]].set_visible(False)
    ax2.tick_params(axis="y", left=False, labelleft=False)
    fig.subplots_adjust(left=.22, right=.98, bottom=.14, top=.96, wspace=.12)
    path = FIG_DIR / "Figure_14_supplier_groups_subgroups.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def save_supplier_cons2000_figure(joined: pd.DataFrame) -> Path:
    data = table_cons2000_delta(joined, ["supplier_type"]).rename(columns={"Supplier Type": "Supplier"})
    data["Supplier"] = pd.Categorical(data["Supplier"], SUPPLIER_ORDER, ordered=True)
    data = data.sort_values("Supplier")
    x = np.arange(len(data))
    width = 0.36
    fig, ax = plt.subplots(figsize=CHART_SIZE_IN)
    fig.patch.set_facecolor("white")
    ax.bar(x - width / 2, data["Baseline total liability (ha)"].to_numpy(dtype=float), width, label="With 2000 rule", color="#AEB5B9")
    ax.bar(x + width / 2, data["Without 2000 total liability (ha)"].to_numpy(dtype=float), width, label="Without 2000 rule", color="#D8743C")
    ax.set_xticks(x)
    ax.set_xticklabels(wrap_axis_labels(data["Supplier"].astype(str), width=14), rotation=0, ha="center", fontsize=FIG_FONT_SIZE)
    polish_chart(fig, ax, "Total liability (ha)")
    add_chart_title(ax, "Figure 15: 2000-rule effect by supplier tier\nTotal liability under the two scenarios")
    add_bar_labels(ax)
    place_legend_right(ax)
    square_chart_layout(fig)
    path = FIG_DIR / "Figure_15_supplier_cons2000_scenario.png"
    fig.savefig(path, dpi=FIG_DPI, facecolor="white")
    plt.close(fig)
    return path


def build_tables_and_figures() -> tuple[dict[str, pd.DataFrame], list[Path], Path]:
    fcm_gta.build_final_workbook()
    secondary_vegetation_by_property(force=True)
    secondary = secondary_vegetation_by_property()
    fc = compliance.add_secondary_vegetation_scenarios(prep_fc(pd.read_parquet(today_path("forest_code_mt_priority_consolidated", "parquet"))), secondary)
    joined = compliance.add_secondary_vegetation_scenarios(prep_fc(pd.read_parquet(today_path("forest_code_gta_final_mt", "parquet"))), secondary)
    fc = apply_input_order(fc)
    joined = apply_input_order(joined)
    joined["supplier_type"] = pd.Categorical(joined["supplier_type"], SUPPLIER_ORDER, ordered=True)
    fc.to_parquet(today_path("forest_code_mt_priority_consolidated_with_secondary", "parquet"), index=False)
    fc.to_csv(today_path("forest_code_mt_priority_consolidated_with_secondary", "csv"), index=False)
    joined.to_parquet(today_path("forest_code_gta_final_mt_with_secondary", "parquet"), index=False)

    active = fc[fc["car_valid"]].copy()
    tables: dict[str, pd.DataFrame] = {}
    tables["Table_01_Database_Sources"] = table_database_sources()
    tables["Table_02_Status"] = table_status(fc)
    tables["Table_03_Active_Status"] = table_status_detail(active)
    tables["Table_04_Size"] = table_size(fc)
    tables["Table_05_Status_Size"] = table_status_by_size(fc)
    tables["Table_06_Vegetation_Cover"] = pd.DataFrame([{
        "Forest area (ha)": active["radam_forest_ha"].sum(),
        "Cerrado area (ha)": active["radam_cerrado_ha"].sum(),
        "Total native vegetation basis (ha)": active["radam_forest_ha"].sum() + active["radam_cerrado_ha"].sum(),
    }])
    tables["Table_07_LR_Deficit_Size"] = (
        active.groupby("size_class", observed=False)
        .agg(**{"Mean LR deficit (ha)": ("rl_adj_deficit_ha", lambda s: s[s > 0].mean() if (s > 0).any() else 0), "Median LR deficit (ha)": ("rl_adj_deficit_ha", lambda s: s[s > 0].median() if (s > 0).any() else 0), "Total LR deficit (ha)": ("rl_adj_deficit_ha", "sum")})
        .reset_index().rename(columns={"size_class": "Size Class"})
    )
    tables["Table_08_LR_Compliance"] = table_lr(active, ["size_class"]).rename(columns={"size_class": "Size Class"})
    tables["Table_09_APP"] = table_app(active, ["size_class"]).rename(columns={"size_class": "Size Class"})
    tables["Table_10_Combined_Size"] = table_combined(active, ["size_class"]).rename(columns={"size_class": "Size Class"})
    tables["Table_11_Mean_Total_Def"] = table_mean_deficit(active)
    tables["Table_12_Cons2000_Overall"] = table_cons2000_compliance(active)
    tables["Table_13_Cons2000_Delta"] = table_cons2000_delta(active)
    tables["Table_14_Cons2000_Size"] = table_cons2000_compliance(active, ["size_class"])
    tables["Table_15_Cons2000_Size_Delta"] = table_cons2000_delta(active, ["size_class"])
    tables["Table_16_SecVeg_Input"] = table_secondary_vegetation_impact(active, ["input_file_type"])
    tables["Table_17_SecVeg_Size"] = table_secondary_vegetation_impact(active, ["input_file_type", "size_class"])
    tables["Table_18_Top20_Mun"] = table_top_municipalities(active)
    tables["Table_19_All_Size"] = table_size_category(fc, "All Properties")
    tables["Table_20_All_LR"] = table_lr_category(fc, "All Properties")
    tables["Table_21_All_APP"] = table_app_category(fc, "All Properties")
    tables["Table_22_All_Combined"] = table_combined_category(fc, "All Properties")
    tables["Table_23_Active_Size"] = table_size_category(active, "Active Properties")
    tables["Table_24_Active_LR"] = table_lr_category(active, "Active Properties")
    tables["Table_25_Active_APP"] = table_app_category(active, "Active Properties")
    tables["Table_26_Active_Combined"] = table_combined_category(active, "Active Properties")
    tables["Table_27_Cattle_Char"] = table_cattle_characterization(joined)
    tables["Table_28_Cattle_Subgroups"] = table_cattle_supplier_subgroups(joined)
    tables["Table_29_Cattle_Combined"] = table_combined(joined, ["supplier_type", "size_class"]).rename(columns={"supplier_type": "Category", "size_class": "Size Class"})
    tables["Table_30_Cattle_Cons2000"] = table_cons2000_compliance(joined, ["supplier_type"])
    tables["Table_31_Cattle_Cons2000_Delta"] = table_cons2000_delta(joined, ["supplier_type"])
    tables["Table_32_Cattle_SecVeg"] = table_secondary_vegetation_impact(joined, ["supplier_type"])
    tables["Table_33_Cattle_LR"] = table_lr(joined, ["supplier_type", "size_class"]).rename(columns={"supplier_type": "Category", "size_class": "Size Class"})
    tables["Table_34_Cattle_Net"] = table_net_liabilities(joined)
    tables["Table_35_Cattle_Size"] = table_cattle_size(joined)
    tables["Table_36_Cattle_LR"] = table_cattle_metric(joined, table_lr_category)
    tables["Table_37_Cattle_APP"] = table_cattle_metric(joined, table_app_category)
    tables["Table_38_Cattle_Combined"] = table_cattle_metric(joined, table_combined_category)
    tables["Table_39_Special_Areas"] = table_special_areas(fc)
    tables["Table_40_Special_Size"] = table_special_by_size(fc)
    tables["Table_41_Special_LR"] = table_special_metric(fc, table_lr_category)
    tables["Table_42_Special_APP"] = table_special_metric(fc, table_app_category)
    tables["Table_43_Special_Combined"] = table_special_metric(fc, table_combined_category)
    tables["Table_44_Mun_Restoration"] = table_top_municipalities(active.assign(app_restore_ha=0), n=20)

    lr_fig = tables["Table_08_LR_Compliance"].copy()
    app_fig = tables["Table_09_APP"].copy()
    comb_fig = tables["Table_10_Combined_Size"].copy()
    mean_fig = tables["Table_11_Mean_Total_Def"].copy()

    figure_paths = []
    study_map = save_study_area_map()
    if study_map is not None:
        figure_paths.append(study_map)
    figure_paths.extend([
        save_stacked_bar(lr_fig, "Size Class", ["To Restore", "To Compensate"], "Figure 4: Legal reserve deficit breakdown", FIG_DIR / "Figure_04_LR_deficit_breakdown.png"),
        save_grouped_bar(
            pd.DataFrame({
                "Biome": ["Forest", "Forest", "Cerrado", "Cerrado"],
                "Metric": ["Adjusted deficit", "Surplus", "Adjusted deficit", "Surplus"],
                "Area (ha)": [
                    active["rl_adj_deficit_forest_ha"].sum(), active["rl_surplus_forest_ha"].sum(),
                    active["rl_adj_deficit_cerrado_ha"].sum(), active["rl_surplus_cerrado_ha"].sum(),
                ],
            }),
            "Biome", "Area (ha)", "Metric", "Figure 5: LR compensation demand and supply by biome", FIG_DIR / "Figure_05_LR_biome_supply_demand.png",
        ),
        save_pie_counts(
            pd.Series([
                ((active["rl_adj_deficit_ha"] + active["app_restore_ha"]) == 0).sum(),
                ((active["rl_adj_deficit_ha"] + active["app_restore_ha"]) > 0).sum(),
            ]),
            pd.Series(["Fully compliant", "Any liability"]),
            "Figure 6: Full compliance status",
            FIG_DIR / "Figure_06_compliance_status.png",
        ),
        save_stacked_bar(comb_fig, "Size Class", ["To Restore (ha)", "To Compensate (ha)"], "Figure 7: Liability for restoration and compensation", FIG_DIR / "Figure_07_restoration_compensation.png"),
        save_grouped_bar(mean_fig.assign(Metric="Mean total deficit"), "Size Class", "Total Mean Deficit (ha)", "Metric", "Figure 8: Mean deficit per property by size class", FIG_DIR / "Figure_08_mean_deficit_size.png"),
    ])
    figure_paths.extend(make_pdf_style_maps(active))
    figure_paths.append(save_vegetation_cover_figure(active))
    figure_paths.append(save_secondary_impact_figure(active))
    figure_paths.append(save_cons2000_overall_figure(active))
    figure_paths.append(save_cons2000_size_figure(active))
    figure_paths.append(save_supplier_subgroups_figure(joined))
    figure_paths.append(save_supplier_cons2000_figure(joined))
    figure_paths.extend(input_audit.write_outputs(input_audit.load_inventory()))

    workbook = OUT_DIR / f"masson_style_final_tables_figures_{kernel.DATE}.xlsx"
    return tables, figure_paths, workbook


def write_workbook(tables: dict[str, pd.DataFrame], figures: list[Path], workbook: Path) -> Path:
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        pd.DataFrame({
            "item": ["reference", "note", "figure_dir"],
            "value": [
                "MS_Masson_Final.pdf table/figure structure",
                "Based on the priority order SIMCAR validated > SIMCAR digital > SIMCAR proxy and the binary GTA supplier rule.",
                str(FIG_DIR),
            ],
        }).to_excel(writer, sheet_name="README", index=False)
        for name, table in tables.items():
            align_table_language(table).rename(columns=display_header).to_excel(writer, sheet_name=name[:31], index=False)
        pd.DataFrame({"figure": [p.name for p in figures], "path": [str(p) for p in figures]}).to_excel(writer, sheet_name="Figures_Index", index=False)

    wb = load_workbook(workbook)
    format_workbook_tables(wb)
    ws = wb["Figures_Index"]
    row = 2
    for fig in figures:
        img = XLImage(str(fig))
        with PILImage.open(fig) as source_image:
            aspect = source_image.height / source_image.width
        img.width = 720
        img.height = int(720 * aspect)
        ws.add_image(img, f"D{row}")
        row += max(24, int(img.height / 20) + 2)
    wb.save(workbook)
    return workbook


def format_workbook_tables(wb: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="FFFFFF")
    total_fill = PatternFill("solid", fgColor="F3F3F3")
    alt_fill = PatternFill("solid", fgColor="FFFFFF")
    body_border = Border(
        left=Side(style=None),
        right=Side(style=None),
        top=Side(style=None),
        bottom=Side(style="hair", color="D9D9D9"),
    )
    header_border = Border(
        left=Side(style=None),
        right=Side(style=None),
        top=Side(style="medium", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    total_border = Border(
        left=Side(style=None),
        right=Side(style=None),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="medium", color="000000"),
    )
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.border = body_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.font = Font(name="Arial", size=TABLE_BODY_FONT_SIZE)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(name="Arial", size=TABLE_HEADER_FONT_SIZE, bold=True)
            cell.border = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_idx in range(2, ws.max_row + 1):
            first_vals = " ".join(str(ws.cell(row_idx, col_idx).value or "").lower() for col_idx in range(1, min(ws.max_column, 2) + 1))
            is_total = "total" in first_vals or "subtotal" in first_vals
            fill = total_fill if is_total else (alt_fill if row_idx % 2 == 0 else None)
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row_idx, col_idx)
                if fill:
                    cell.fill = fill
                if is_total:
                    cell.font = Font(name="Arial", size=TABLE_BODY_FONT_SIZE, bold=True)
                    cell.border = total_border
                if isinstance(cell.value, (int, float)):
                    header = str(ws.cell(1, col_idx).value or "")
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '0.0"%"' if header.startswith("%") else "#,##0.0" if abs(float(cell.value)) < 1000 and not float(cell.value).is_integer() else "#,##0"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            header = str(ws.cell(1, col_idx).value or "")
            max_len = len(header)
            for row_idx in range(2, min(ws.max_row, 80) + 1):
                max_len = max(max_len, len(str(ws.cell(row_idx, col_idx).value or "")))
            if header in {"Path", "Notes", "Role in final product"}:
                width = min(max(max_len * 0.8, 18), 48)
            elif header in {"Category", "Size Class", "Municipality", "Database", "Registration Status", "Special Area"}:
                width = min(max(max_len * 0.9, 14), 26)
            else:
                width = min(max(max_len * 0.85, 10), 18)
            ws.column_dimensions[letter].width = width
        ws.row_dimensions[1].height = 36


def main() -> Path:
    tables, figures, workbook = build_tables_and_figures()
    out = write_workbook(tables, figures, workbook)
    print("Masson-style workbook:", out)
    print("Tables:", len(tables))
    print("Figures:", len(figures))
    for fig in figures:
        print(fig)
    return out


if __name__ == "__main__":
    main()

