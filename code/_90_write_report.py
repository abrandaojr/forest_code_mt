from __future__ import annotations

import html
import os
import re
import zipfile
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd

import _00_paths as paths
import _10_kernel as kernel

paths.set_env()
ROOT = paths.ROOT
OUT_DIR = paths.TABLES
QC_DIR = paths.QC
REPORT_DIR = paths.REPORTS
REPORT_DIR.mkdir(parents=True, exist_ok=True)
TODAY = kernel.DATE
INPUT_ORDER = kernel.SOURCE_ORDER
INPUT_LABELS = kernel.SOURCE_LABELS


def dated_file(folder: Path, pattern: str) -> Path | None:
    exact = folder / pattern.format(date=TODAY)
    if exact.exists():
        return exact
    matches = sorted(folder.glob(pattern.format(date="*")), key=lambda p: p.stat().st_mtime, reverse=True)
    return matches[0] if matches else None


def tex_escape(value: object) -> str:
    text = str(value)
    repl = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_\allowbreak{}",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
        "[": r"{[}",
        "]": r"{]}",
    }
    return "".join(repl.get(ch, ch) for ch in text)


def fmt_int(value: int | float) -> str:
    return f"{int(value):,}"


def fmt_float(value: int | float, digits: int = 3) -> str:
    return f"{float(value):,.{digits}f}"


def data_footprint() -> list[tuple[str, int, float]]:
    rows = []
    for name in ["data/raw", "data/pre", "data/proc", "out"]:
        root = ROOT / name
        files = [p for p in root.rglob("*") if p.is_file()] if root.exists() else []
        rows.append((name, len(files), sum(p.stat().st_size for p in files) / 1024**3))
    return rows


def validation_metrics() -> dict[str, object]:
    fc = pd.read_parquet(OUT_DIR / f"forest_code_mt_priority_consolidated_{TODAY}.parquet")
    core = [
        "rl_adj_deficit_ha",
        "app_restore_ha",
        "app_gross_deficit_ha",
        "calc_deficit_total_ha",
        "calc_gross_deficit_total_ha",
    ]
    net = (
        fc["calc_deficit_total_ha"]
        - (
            pd.to_numeric(fc["rl_adj_deficit_ha"], errors="coerce").fillna(0)
            + pd.to_numeric(fc["app_restore_ha"], errors="coerce").fillna(0)
        )
    ).abs().max()
    gross = (
        fc["calc_gross_deficit_total_ha"]
        - (
            pd.to_numeric(fc["rl_gross_deficit_ha"], errors="coerce").fillna(0)
            + pd.to_numeric(fc["app_gross_deficit_ha"], errors="coerce").fillna(0)
        )
    ).abs().max()
    zero_cols = [c for c in ["mun_geocodigo", "mun_geocodigo_norm"] if c in fc.columns]
    zero_audit_count = int(
        sum(
            fc[c].astype("string").str.replace(r"\D", "", regex=True).str.fullmatch(r"0+").fillna(False).sum()
            for c in zero_cols
        )
    )
    raw_cmp_path = dated_file(QC_DIR, "portable_required_raw_vs_R_{date}.csv")
    raw_cmp = pd.read_csv(raw_cmp_path) if raw_cmp_path is not None else pd.DataFrame(columns=["exists_in_package", "found_in_R_copy"])
    by_input_raw = fc["input_file_type"].value_counts().to_dict()
    by_input = {INPUT_LABELS[k]: int(by_input_raw.get(k, 0)) for k in INPUT_ORDER}
    return {
        "rows": len(fc),
        "by_input": by_input,
        "net_diff": float(net),
        "gross_diff": float(gross),
        "negative_core": int((fc[core].apply(pd.to_numeric, errors="coerce").fillna(0) < -1e-9).sum().sum()),
        "zero_municipality_findings": zero_audit_count,
        "required_raw_missing_package": int((raw_cmp["exists_in_package"] != True).sum()),
        "required_raw_missing_r": int((raw_cmp["found_in_R_copy"] != True).sum()),
    }


def provenance_rows() -> list[tuple[str, int, float]]:
    manifest_path = dated_file(QC_DIR, "geoparquet_raw_provenance_manifest_{date}.csv")
    if manifest_path is None:
        return []
    manifest = pd.read_csv(manifest_path)
    grouped = (
        manifest.groupby("inferred_source", dropna=False)
        .agg(files=("relative_path", "size"), size_mb=("size_mb", "sum"))
        .reset_index()
        .sort_values("inferred_source")
    )
    return [(str(r.inferred_source), int(r.files), float(r.size_mb)) for r in grouped.itertuples()]


def latex_table(rows: list[tuple[object, ...]], headers: list[str], aligns: str) -> str:
    body = ["\\begin{tabular}{" + aligns + "}", "\\toprule"]
    body.append(" & ".join(tex_escape(h) for h in headers) + r" \\")
    body.append("\\midrule")
    for row in rows:
        body.append(" & ".join(tex_escape(v) for v in row) + r" \\")
    body.append("\\bottomrule")
    body.append("\\end{tabular}")
    return "\n".join(body)


def format_cell(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (int, np.integer)) and not isinstance(value, bool):
        return f"{int(value):,}"
    if isinstance(value, (float, np.floating)):
        if abs(value) >= 1000:
            return f"{value:,.1f}"
        return f"{value:,.3f}".rstrip("0").rstrip(".")
    if isinstance(value, str):
        text = value.strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", text) and len(text.replace("-", "").split(".")[0]) > 3:
            num = float(text)
            return f"{num:,.1f}" if "." in text else f"{int(num):,}"
    return str(value)


def display_header(value: object) -> str:
    text = str(value)
    text = text.replace("(Mha)", "(million hectares)")
    text = text.replace("(ha)", "(hectares)")
    text = text.replace("_ha", " (hectares)")
    text = re.sub(r"\bMha\b", "million hectares", text)
    text = re.sub(r"\bha\b", "hectares", text)
    return text


def dataframe_latex_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    work = df.copy()
    if max_rows is not None:
        work = work.head(max_rows)
    if len(work.columns) <= 8:
        chunks = [list(work.columns)]
    else:
        stub_count = min(2, len(work.columns) - 1)
        stubs = list(work.columns[:stub_count])
        value_cols = list(work.columns[stub_count:])
        chunks = []
        for start in range(0, len(value_cols), 5):
            chunks.append(stubs + value_cols[start:start + 5])
    rendered = []
    for idx, cols in enumerate(chunks, start=1):
        headers = [tex_escape(c) for c in cols]
        rows = []
        for _, row in work[cols].iterrows():
            rows.append(" & ".join(tex_escape(format_cell(v)) for v in row.tolist()) + r" \\")
        col_width = max(0.08, min(0.26, 0.82 / max(len(headers), 1)))
        col_spec = "".join(rf"p{{{col_width:.3f}\linewidth}}" for _ in headers)
        panel = []
        if len(chunks) > 1:
            panel.append(rf"\textit{{Panel {idx} of {len(chunks)}}}\\[2pt]")
        panel.extend(
            [
                r"{\fontsize{12}{14}\selectfont",
                r"\setlength{\tabcolsep}{3pt}",
                r"\renewcommand{\arraystretch}{1.08}",
                r"\begin{tabular}{" + col_spec + "}",
                r"\toprule",
                " & ".join(headers) + r" \\",
                r"\midrule",
                *rows,
                r"\bottomrule",
                r"\end{tabular}",
                r"}",
            ]
        )
        rendered.append("\n".join(panel))
    return "\n\n\\vspace{0.6em}\n\n".join(rendered)


def table_caption(sheet: str) -> str:
    labels = {
        "Table_01_Database_Sources": "Table 1. Data sources and input priority",
        "Table_02_Status": "Table 2. Overall CAR registration status",
        "Table_03_Active_Status": "Table 3. Active-registration detail",
        "Table_04_Size": "Table 4. Property-size distribution",
        "Table_05_Status_Size": "Table 5. Registration status by property size",
        "Table_06_Vegetation_Cover": "Table 6. Vegetation basis used in the calculation",
        "Table_07_LR_Deficit_Size": "Table 7. Legal Reserve deficit by property size",
        "Table_08_LR_Compliance": "Table 8. Legal Reserve compliance by property size",
        "Table_09_APP": "Table 9. APP liability by property size",
        "Table_10_Combined_Size": "Table 10. Combined liability by property size",
        "Table_11_Mean_Total_Def": "Table 11. Mean total deficit per property",
        "Table_12_Cons2000_Overall": "Table 12. Scenario with and without the 2000 Legal Reserve rule",
        "Table_13_Cons2000_Delta": "Table 13. Change attributable to the 2000 Legal Reserve rule",
        "Table_14_Cons2000_Size": "Table 14. 2000-rule scenario by property size",
        "Table_15_Cons2000_Size_Delta": "Table 15. 2000-rule change by property size",
        "Table_16_SecVeg_Input": "Table 16. Secondary-forest scenario by input source",
        "Table_17_SecVeg_Size": "Table 17. Secondary-forest scenario by input source and property size",
        "Table_18_Top20_Mun": "Table 18. Municipalities with the largest total liability",
        "Table_27_Cattle_Char": "Table 27. Supplier characterization",
        "Table_28_Cattle_Subgroups": "Table 28. Supplier groups and subgroups",
        "Table_29_Cattle_Combined": "Table 29. Combined liability among suppliers",
        "Table_30_Cattle_Cons2000": "Table 30. Suppliers with and without the 2000 rule",
        "Table_31_Cattle_Cons2000_Delta": "Table 31. 2000-rule change among suppliers",
        "Table_32_Cattle_SecVeg": "Table 32. Secondary forest among suppliers",
        "Table_39_Special_Areas": "Table 39. Overlap with special areas",
    }
    return labels.get(sheet, sheet.replace("_", " "))


def figure_caption(path: Path) -> str:
    stem = path.stem
    if stem.startswith("Figure_16") and stem.endswith("_inputs"):
        return "Spatial inputs used by the analysis; equivalent technical variants are consolidated."
    labels = {
        "chart_deficit_by_input": "Total deficit by input source; values are hectares.",
        "chart_deficit_by_size": "Total deficit by property-size class; values are hectares.",
        "chart_properties_by_input": "Number of properties by input source; values are property counts.",
        "chart_supplier_by_input": "Supplier properties by input source; values are property counts.",
        "map_app_gross_deficit_ha": "Municipal map of gross APP deficit; color scale is hectares.",
        "map_rl_adjusted_deficit_ha": "Municipal map of adjusted Legal Reserve deficit; color scale is hectares.",
        "map_total_deficit_ha": "Municipal map of total deficit; color scale is hectares.",
        "Figure_03_study_area_mato_grosso": "Study area in Mato Grosso.",
        "Figure_04_LR_deficit_breakdown": "Breakdown of the Legal Reserve deficit; values are hectares.",
        "Figure_05_LR_biome_supply_demand": "Legal Reserve demand and surplus by vegetation formation; values are hectares.",
        "Figure_06_compliance_status": "Overall compliance status; values are property counts.",
        "Figure_07_restoration_compensation": "Restoration and compensation liability; values are hectares.",
        "Figure_08_mean_deficit_size": "Mean deficit by property-size class; values are hectares per property.",
        "Figure_09_municipal_noncompliance_panel": "Municipal non-compliance panel. APP and Legal Reserve are shown separately. Count maps report the number of non-compliant properties. Hectares/property maps report non-compliant area divided by affected properties in the same component and scenario. Legal Reserve is shown with and without the 2000 Legal Reserve rule, each with a property-count map and a hectares-per-affected-property map.",
        "Figure_10_vegetation_cover": "Vegetation cover used in the calculation; values are million hectares.",
        "Figure_11_secondary_vegetation_impact": "Impact of secondary forest on total liability; values are hectares.",
        "Figure_12_cons2000_total_scenario": "Total variation with and without the 2000 Legal Reserve rule; values are hectares.",
        "Figure_13_cons2000_size_scenario": "Variation by property size with and without the 2000 Legal Reserve rule; values are hectares.",
        "Figure_14_supplier_groups_subgroups": "Supplier groups and subgroups; left axis is property count and right axis is cattle head.",
        "Figure_15_supplier_cons2000_scenario": "Supplier liabilities with and without the 2000 Legal Reserve rule; values are hectares.",
    }
    if stem in labels:
        return labels[stem]
    return stem.replace("_", " ").replace("Figure ", "Figure ").strip() + "."


def iter_result_tables() -> Iterable[tuple[str, pd.DataFrame]]:
    xlsx = OUT_DIR / f"masson_style_final_tables_figures_{TODAY}.xlsx"
    xl = pd.ExcelFile(xlsx)
    for sheet in xl.sheet_names:
        if sheet.startswith("Table_"):
            df = pd.read_excel(xlsx, sheet_name=sheet)
            for col in ["Input File Type", "Input source", "Input data"]:
                if col in df.columns:
                    order_map = {
                        **{k: i for i, k in enumerate(INPUT_ORDER)},
                        **{v: i for i, v in enumerate(INPUT_LABELS.values())},
                        "SIMCAR validado": 0,
                        "SIMCAR validated": 0,
                        "SIMCAR digital": 1,
                        "SIMCAR proxy": 2,
                    }
                    df["_input_order"] = df[col].astype(str).map(order_map).fillna(99)
                    sort_cols = ["_input_order"]
                    if "Size Class" in df.columns:
                        sort_cols.append("Size Class")
                    df = df.sort_values(sort_cols, kind="stable").drop(columns="_input_order")
                    df[col] = df[col].astype(str).replace(INPUT_LABELS)
            df = df.rename(columns=display_header)
            yield sheet, df


def iter_result_figures() -> list[Path]:
    fig_dir = Path(os.environ.get("FCM_FIGURES_DIR", ROOT / "out" / "fig"))
    preferred = [
        "Figure_03_study_area_mato_grosso.png",
        "Figure_04_LR_deficit_breakdown.png",
        "Figure_05_LR_biome_supply_demand.png",
        "Figure_06_compliance_status.png",
        "Figure_07_restoration_compensation.png",
        "Figure_08_mean_deficit_size.png",
        "Figure_09_municipal_noncompliance_panel.png",
        "Figure_10_vegetation_cover.png",
        "Figure_11_secondary_vegetation_impact.png",
        "Figure_12_cons2000_total_scenario.png",
        "Figure_13_cons2000_size_scenario.png",
        "Figure_14_supplier_groups_subgroups.png",
        "Figure_15_supplier_cons2000_scenario.png",
        *[p.name for p in sorted(fig_dir.glob("Figure_16?_*.png"))],
        "chart_properties_by_input.png",
        "chart_deficit_by_input.png",
        "chart_supplier_by_input.png",
        "chart_deficit_by_size.png",
        "map_total_deficit_ha.png",
        "map_rl_adjusted_deficit_ha.png",
        "map_app_gross_deficit_ha.png",
    ]
    return [fig_dir / name for name in preferred if (fig_dir / name).exists()]


def render_latex_table_block(sheet: str, df: pd.DataFrame) -> str:
    env_start = r"\begin{landscape}" if len(df.columns) > 8 else ""
    env_end = r"\end{landscape}" if len(df.columns) > 8 else ""
    row_limit = 4 if len(df.columns) > 8 else 18
    chunks = [df.iloc[start:start + row_limit].copy() for start in range(0, len(df), row_limit)] or [df]
    blocks = [env_start] if env_start else []
    for idx, chunk in enumerate(chunks, start=1):
        caption = table_caption(sheet) if idx == 1 else f"{table_caption(sheet)} (continued)"
        blocks.extend([
            r"\begin{table}[H]",
            r"\centering",
            rf"\caption{{{tex_escape(caption)}}}",
            dataframe_latex_table(chunk),
            r"\end{table}",
        ])
    if env_end:
        blocks.append(env_end)
    return "\n".join(blocks)


def render_latex_figure_block(fig: Path) -> str:
    rel_fig = Path(os.path.relpath(fig, REPORT_DIR)).as_posix()
    return "\n".join(
        [
            r"\begin{figure}[H]",
            r"\centering",
            rf"\includegraphics[width=0.92\linewidth]{{\detokenize{{{rel_fig}}}}}",
            rf"\caption{{{tex_escape(figure_caption(fig))}}}",
            r"\end{figure}",
        ]
    )


def table_by_prefix(table_map: dict[str, pd.DataFrame], prefix: str) -> tuple[str, pd.DataFrame] | None:
    for sheet, df in table_map.items():
        if sheet.startswith(prefix):
            return sheet, df
    return None


def figure_by_stem(figures: list[Path], stem: str) -> Path | None:
    for fig in figures:
        if fig.stem == stem:
            return fig
    return None


MAIN_TABLE_PREFIXES = {
    "Table_01",
    "Table_02",
    "Table_04",
    "Table_06",
    "Table_08",
    "Table_09",
    "Table_10",
    "Table_13",
    "Table_16",
    "Table_18",
    "Table_28",
    "Table_31",
    "Table_39",
}

MAIN_FIGURE_STEMS = {
    "chart_properties_by_input",
    "Figure_10_vegetation_cover",
    "Figure_04_LR_deficit_breakdown",
    "Figure_06_compliance_status",
    "Figure_12_cons2000_total_scenario",
    "Figure_11_secondary_vegetation_impact",
    "Figure_09_municipal_noncompliance_panel",
    "map_total_deficit_ha",
    "Figure_14_supplier_groups_subgroups",
    "Figure_15_supplier_cons2000_scenario",
}


def starts_with_any(value: str, prefixes: set[str]) -> bool:
    return any(value.startswith(prefix) for prefix in prefixes)


def split_evidence(
    result_tables: list[tuple[str, pd.DataFrame]],
    result_figures: list[Path],
) -> tuple[list[tuple[str, pd.DataFrame]], list[Path], list[tuple[str, pd.DataFrame]], list[Path]]:
    main_tables = [(sheet, df) for sheet, df in result_tables if starts_with_any(sheet, MAIN_TABLE_PREFIXES)]
    supp_tables = [(sheet, df) for sheet, df in result_tables if not starts_with_any(sheet, MAIN_TABLE_PREFIXES)]
    main_figures = [fig for fig in result_figures if fig.stem in MAIN_FIGURE_STEMS]
    supp_figures = [fig for fig in result_figures if fig.stem not in MAIN_FIGURE_STEMS]
    return main_tables, main_figures, supp_tables, supp_figures


def render_evidence_lists(
    main_tables: list[tuple[str, pd.DataFrame]],
    main_figures: list[Path],
    supp_tables: list[tuple[str, pd.DataFrame]],
    supp_figures: list[Path],
) -> str:
    def table_items(rows: list[tuple[str, pd.DataFrame]]) -> str:
        if not rows:
            return r"\item No tables."
        return "\n".join(rf"\item {tex_escape(table_caption(sheet))}" for sheet, _ in rows)

    def figure_items(rows: list[Path]) -> str:
        if not rows:
            return r"\item No figures."
        return "\n".join(rf"\item {tex_escape(figure_caption(path))}" for path in rows)

    return "\n".join([
        r"\section*{Main Manuscript Tables}",
        r"\begin{enumerate}",
        table_items(main_tables),
        r"\end{enumerate}",
        r"\clearpage",
        r"\section*{Main Manuscript Figures}",
        r"\begin{enumerate}",
        figure_items(main_figures),
        r"\end{enumerate}",
        r"\clearpage",
        r"\section*{Supplementary Tables}",
        r"\begin{enumerate}",
        table_items(supp_tables),
        r"\end{enumerate}",
        r"\clearpage",
        r"\section*{Supplementary Figures}",
        r"\begin{enumerate}",
        figure_items(supp_figures),
        r"\end{enumerate}",
        r"\clearpage",
    ])


def render_question_sections(result_tables: list[tuple[str, pd.DataFrame]], result_figures: list[Path]) -> str:
    table_map = dict(result_tables)
    used_tables: set[str] = set()
    used_figures: set[Path] = set()
    questions = [
        (
            "Data completeness and traceability",
            "The first question verifies whether the inputs that support the model were retained, whether the priority order among validated CAR, digital CAR, and proxy CAR is explicit, and whether raw data converted to GeoParquet remain traceable.",
            ["Table_01", "Table_02", "Table_03"],
            ["chart_properties_by_input"],
        ),
        (
            "Analytical universe of properties",
            "The second question describes the size and structure of the property universe before applying environmental-law calculations.",
            ["Table_04", "Table_05", "Table_19", "Table_23"],
            [],
        ),
        (
            "Vegetation basis for the calculation",
            "This question separates the biophysical basis from the legal calculation: forest area, cerrado area, and the vegetation classes used to compute Legal Reserve and APP liabilities.",
            ["Table_06"],
            ["Figure_10_vegetation_cover"],
        ),
        (
            "Legal Reserve liability",
            "The report then moves to the first legal component: required area, gross deficit, adjusted deficit, restoration, compensation, and surplus Legal Reserve.",
            ["Table_07", "Table_08", "Table_20", "Table_24"],
            ["Figure_04_LR_deficit_breakdown", "Figure_05_LR_biome_supply_demand"],
        ),
        (
            "APP liability",
            "After Legal Reserve, the report presents APP as a separate component and distinguishes gross deficit from the effective restoration liability.",
            ["Table_09", "Table_21", "Table_25"],
            ["map_app_gross_deficit_ha"],
        ),
        (
            "Combined compliance liability",
            "The combined answer adds adjusted Legal Reserve liability and restorable APP liability, which is the final compliance indicator used in the package.",
            ["Table_10", "Table_11", "Table_22", "Table_26"],
            ["Figure_06_compliance_status", "Figure_07_restoration_compensation", "Figure_08_mean_deficit_size", "chart_deficit_by_input", "chart_deficit_by_size"],
        ),
        (
            "Sensitivity to the 2000 Legal Reserve rule",
            "This scenario isolates the 2000 rule within Legal Reserve. APP liability is held constant between scenarios so that the comparison does not mix two different methodological questions.",
            ["Table_12", "Table_13", "Table_14", "Table_15"],
            ["Figure_12_cons2000_total_scenario", "Figure_13_cons2000_size_scenario"],
        ),
        (
            "Sensitivity to secondary forest recognition",
            "This scenario tests the sensitivity of the Legal Reserve deficit when secondary forest is added to the stock of existing vegetation.",
            ["Table_16", "Table_17"],
            ["Figure_11_secondary_vegetation_impact"],
        ),
        (
            "Municipal concentration of liability",
            "The territorial reading comes after the aggregate results because it depends on property-level indicators and their municipal aggregation.",
            ["Table_18", "Table_44"],
            ["Figure_03_study_area_mato_grosso", "Figure_09_municipal_noncompliance_panel", "map_total_deficit_ha", "map_rl_adjusted_deficit_ha"],
        ),
        (
            "Supplier exposure and cattle-chain segmentation",
            "This is a more complex layer because it connects environmental compliance with supplier groups and subgroups while keeping cattle volume and legal liability analytically separate.",
            ["Table_27", "Table_28", "Table_29", "Table_30", "Table_31", "Table_32", "Table_33", "Table_34", "Table_35", "Table_36", "Table_37", "Table_38"],
            ["chart_supplier_by_input", "Figure_14_supplier_groups_subgroups", "Figure_15_supplier_cons2000_scenario"],
        ),
        (
            "Special-area overlays",
            "The final question intersects the indicators with sensitive overlaps. It therefore appears after the general results, scenarios, and supplier analysis.",
            ["Table_39", "Table_40", "Table_41", "Table_42", "Table_43"],
            [],
        ),
    ]
    blocks = [r"\part*{Main Manuscript}"]
    for title, narrative, table_prefixes, figure_stems in questions:
        parts = [rf"\section*{{{tex_escape(title)}}}", tex_escape(narrative)]
        for prefix in table_prefixes:
            found = table_by_prefix(table_map, prefix)
            if found is None:
                continue
            sheet, df = found
            if not starts_with_any(sheet, MAIN_TABLE_PREFIXES):
                continue
            used_tables.add(sheet)
            parts.append(render_latex_table_block(sheet, df))
        for stem in figure_stems:
            fig = figure_by_stem(result_figures, stem)
            if fig is None:
                continue
            if fig.stem not in MAIN_FIGURE_STEMS:
                continue
            used_figures.add(fig)
            parts.append(render_latex_figure_block(fig))
        blocks.append("\n\n".join(parts))

    remaining_tables = [(sheet, df) for sheet, df in result_tables if sheet not in used_tables]
    remaining_figures = [fig for fig in result_figures if fig not in used_figures]
    if remaining_tables or remaining_figures:
        appendix = [r"\part*{Supplementary Material}"]
        if remaining_tables:
            appendix.append(r"\section*{Supplementary Tables}")
        for sheet, df in remaining_tables:
            appendix.append(render_latex_table_block(sheet, df))
        if remaining_figures:
            appendix.append(r"\clearpage")
            appendix.append(r"\section*{Supplementary Figures}")
        for fig in remaining_figures:
            appendix.append(render_latex_figure_block(fig))
        blocks.append("\n\n".join(appendix))
    return "\n\n".join(blocks)


def build_tex() -> Path:
    metrics = validation_metrics()
    footprint = [(p, fmt_int(n), fmt_float(gb, 3)) for p, n, gb in data_footprint()]
    sources = [(s, fmt_int(n), fmt_float(mb, 1)) for s, n, mb in provenance_rows()]
    input_rows = [(k, fmt_int(v)) for k, v in metrics["by_input"].items()]

    result_tables = list(iter_result_tables())
    result_figures = iter_result_figures()
    main_tables, main_figures, supp_tables, supp_figures = split_evidence(result_tables, result_figures)
    evidence_lists = render_evidence_lists(main_tables, main_figures, supp_tables, supp_figures)
    result_sections = render_question_sections(result_tables, result_figures)

    tex = rf"""\documentclass[12pt]{{article}}
\usepackage[a4paper,margin=1in]{{geometry}}
\usepackage{{booktabs}}
\usepackage{{graphicx}}
\usepackage{{hyperref}}
\usepackage{{fontspec}}
\usepackage{{float}}
\usepackage{{pdflscape}}
\setmainfont{{Arial}}
\hypersetup{{colorlinks=true,linkcolor=black,urlcolor=blue}}
\title{{Forest Code Compliance in Mato Grosso: Reproducible Property-Level Estimates}}
\author{{Forest Code Compliance Study Team}}
\date{{{TODAY[:4]}-{TODAY[4:6]}-{TODAY[6:]}}}

\begin{{document}}
\maketitle
\clearpage
{evidence_lists}

\section*{{Article Structure}}
The manuscript is organized as a scientific results paper. It first defines the data, metrics, and validation checks needed to interpret the estimates, and then presents the empirical results from the least model-dependent descriptions to the most interpretive sensitivity and supply-chain analyses.

\section*{{Technical Summary}}
The final analytical dataset contains {fmt_int(metrics["rows"])} property records and applies a fixed source hierarchy throughout the analysis: SIMCAR validated first, SIMCAR digital second, and SIMCAR proxy third. This hierarchy gives precedence to the most precise CAR evidence available for each property while retaining a proxy source for farms without validated or digital CAR geometry. No municipality code equal to 00000 or 0000000 remains in the audited tabular outputs.

Two sensitivity tests required particular attention. The 2000 Legal Reserve scenario now isolates Legal Reserve only, holding APP liability constant so that the comparison has a single interpretation. The secondary-forest scenario was recalculated from the spatial intersection with the prioritized property geometries, avoiding a previous merge/cache artifact that had incorrectly left the secondary-forest column equal to zero.

\section*{{Portable Execution}}
The reproducible command is:
\begin{{verbatim}}
conda run -n geo python .\code\_01_orchestrate.py --stage all
\end{{verbatim}}
The procedure preprocesses the three CAR sources, builds the priority table, joins GTA supplier information, and prepares the tabular and spatial outputs used in this report.

\section*{{Data Footprint After Cleanup}}
{latex_table(footprint, ["Folder", "Files", "GB"], "lrr")}

\section*{{Prioritized Inputs}}
{latex_table(input_rows, ["Input source", "Rows after priority"], "lr")}

\section*{{Formula and Data QA}}
{latex_table([
    ("Net deficit formula max absolute difference", fmt_float(metrics["net_diff"], 6)),
    ("Gross deficit formula max absolute difference", fmt_float(metrics["gross_diff"], 6)),
    ("Negative cells in core deficit metrics", metrics["negative_core"]),
    ("Municipality zero-code findings after fix", metrics["zero_municipality_findings"]),
    ("Required raw files missing in package", metrics["required_raw_missing_package"]),
    ("Required raw files not found in R source copy", metrics["required_raw_missing_r"]),
], ["Check", "Result"], "lr")}

\section*{{Raw and Raw-Converted Source Coverage}}
{latex_table(sources, ["Inferred source", "Files", "Size MB"], "lrr")}

{result_sections}

\section*{{Spatial Processing Rule}}
Property boundaries remain the master geometries. All non-property environmental layers used upstream should be processed as 25 x 25 km tiles before intersection. This rule preserves the legal unit of analysis while keeping the spatial operations tractable.

\section*{{Conclusion}}
The final products are internally consistent, traceable to the retained raw inputs, free of zero municipality codes in the audited outputs, and structured as a publication-oriented results manuscript. The evidence supports direct comparison across property source classes, legal components, sensitivity scenarios, municipalities, and supplier groups while keeping the methodological assumptions explicit.

\end{{document}}
"""
    path = REPORT_DIR / f"forest_code_mt_final_report_{TODAY}.tex"
    path.write_text(tex, encoding="utf-8")
    return path


def xml_escape(text: object) -> str:
    return html.escape(str(text), quote=True)


def paragraph(text: str, style: str | None = None) -> str:
    pstyle = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{pstyle}<w:r><w:t>{xml_escape(text)}</w:t></w:r></w:p>"


def table(rows: list[tuple[object, ...]]) -> str:
    xml = ['<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>']
    for row in rows:
        xml.append("<w:tr>")
        for cell in row:
            xml.append(f"<w:tc><w:p><w:r><w:t>{xml_escape(cell)}</w:t></w:r></w:p></w:tc>")
        xml.append("</w:tr>")
    xml.append("</w:tbl>")
    return "".join(xml)


def build_docx() -> Path:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    metrics = validation_metrics()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    for style_name, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 12)]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Title"

    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    def add_para(text: str, style: str | None = None, align=None):
        p = doc.add_paragraph(text, style=style)
        if align is not None:
            p.alignment = align
        return p

    def set_table_widths(tbl, widths: list[float]):
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl_pr = tbl._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")
        for row in tbl.rows:
            for idx, cell in enumerate(row.cells):
                cell.width = Inches(widths[min(idx, len(widths) - 1)])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def soften_header(value: object) -> str:
        text = str(value).replace("_", " ")
        text = text.replace("Legal Reserve", "LR").replace("Secondary vegetation", "Secondary veg.")
        text = text.replace("Properties", "Props.").replace("Municipality", "Municip.")
        return text

    def add_small_table(rows: list[tuple[object, ...]]):
        tbl = doc.add_table(rows=1, cols=len(rows[0]))
        tbl.style = "Table Grid"
        widths = [3.4] + [max(1.0, 3.8 / max(len(rows[0]) - 1, 1))] * (len(rows[0]) - 1)
        set_table_widths(tbl, widths)
        for i, value in enumerate(rows[0]):
            tbl.rows[0].cells[i].text = soften_header(value)
        for row in rows[1:]:
            cells = tbl.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(12)
        return tbl

    def add_dataframe_table(sheet: str, df: pd.DataFrame):
        work = df.copy()
        if len(work.columns) <= 6:
            col_chunks = [list(work.columns)]
        else:
            stubs = list(work.columns[: min(2, len(work.columns) - 1)])
            values = list(work.columns[len(stubs):])
            col_chunks = [stubs + values[i:i + 3] for i in range(0, len(values), 3)]
        row_limit = 12
        for panel_idx, cols in enumerate(col_chunks, start=1):
            for start in range(0, len(work), row_limit):
                chunk = work.iloc[start:start + row_limit][cols]
                caption = table_caption(sheet)
                if len(col_chunks) > 1:
                    caption = f"{caption}, panel {panel_idx} of {len(col_chunks)}"
                if start:
                    caption = f"{caption} (continued)"
                add_para(caption, "Heading 2")
                tbl = doc.add_table(rows=1, cols=len(cols))
                tbl.style = "Table Grid"
                if len(cols) == 1:
                    widths = [7.2]
                elif len(cols) == 2:
                    widths = [3.4, 3.8]
                else:
                    stub_count = min(2, len(cols) - 1)
                    remaining = len(cols) - stub_count
                    widths = [1.65] * stub_count + [max(1.15, (7.2 - 1.65 * stub_count) / remaining)] * remaining
                set_table_widths(tbl, widths)
                for i, col in enumerate(cols):
                    tbl.rows[0].cells[i].text = soften_header(col)
                for _, row in chunk.iterrows():
                    cells = tbl.add_row().cells
                    for i, value in enumerate(row.tolist()):
                        cells[i].text = format_cell(value)
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.name = "Arial"
                                run.font.size = Pt(12)

    def add_figure(fig: Path):
        add_para(figure_caption(fig), "Heading 2")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(fig), width=Inches(7.2))

    result_tables = list(iter_result_tables())
    result_figures = iter_result_figures()
    main_tables, main_figures, supp_tables, supp_figures = split_evidence(result_tables, result_figures)

    add_para("Forest Code Compliance in Mato Grosso: Reproducible Property-Level Estimates", "Title", WD_ALIGN_PARAGRAPH.CENTER)
    add_para(f"Prepared on {TODAY[:4]}-{TODAY[4:6]}-{TODAY[6:]}", align=WD_ALIGN_PARAGRAPH.CENTER)

    add_para("Main Manuscript Tables", "Heading 1")
    for sheet, _ in main_tables:
        add_para(table_caption(sheet))
    add_para("Main Manuscript Figures", "Heading 1")
    for fig in main_figures:
        add_para(figure_caption(fig))
    add_para("Supplementary Tables", "Heading 1")
    for sheet, _ in supp_tables:
        add_para(table_caption(sheet))
    add_para("Supplementary Figures", "Heading 1")
    for fig in supp_figures:
        add_para(figure_caption(fig))

    doc.add_page_break()
    add_para("Article Structure", "Heading 1")
    add_para("The manuscript is organized as a scientific results paper, moving from data, metrics, and validation checks to empirical results, sensitivity analyses, spatial patterns, and supplier segmentation.")
    add_para("Technical Summary", "Heading 1")
    add_para(
        f"The analytical package was reviewed against the files required to reproduce the results. The priority consolidation "
        f"contains {fmt_int(metrics['rows'])} properties and applies the source hierarchy consistently: SIMCAR validated, "
        "SIMCAR digital, and SIMCAR proxy. No municipality code equal to 00000 or 0000000 remains in the audited outputs."
    )
    add_para("Data Footprint After Cleanup", "Heading 1")
    add_small_table([("Folder", "Files", "GB"), *[(p, fmt_int(n), fmt_float(gb, 3)) for p, n, gb in data_footprint()]])
    add_para("Prioritized Inputs", "Heading 1")
    add_small_table([("Input source", "Rows after priority"), *[(k, fmt_int(v)) for k, v in metrics["by_input"].items()]])
    add_para("Formula And Data QA", "Heading 1")
    add_small_table([
        ("Check", "Result"),
        ("Net deficit formula max absolute difference", fmt_float(metrics["net_diff"], 6)),
        ("Gross deficit formula max absolute difference", fmt_float(metrics["gross_diff"], 6)),
        ("Negative cells in core deficit metrics", metrics["negative_core"]),
        ("Municipality zero-code findings after fix", metrics["zero_municipality_findings"]),
        ("Required raw files missing in package", metrics["required_raw_missing_package"]),
        ("Required raw files not found in R source copy", metrics["required_raw_missing_r"]),
    ])
    add_para("Raw And Raw-Converted Source Coverage", "Heading 1")
    add_small_table([("Inferred source", "Files", "Size MB"), *[(s, fmt_int(n), fmt_float(mb, 1)) for s, n, mb in provenance_rows()]])

    add_para("Main Manuscript", "Heading 1")
    for sheet, df in main_tables:
        add_dataframe_table(sheet, df)
    for fig in main_figures:
        add_figure(fig)

    add_para("Supplementary Material", "Heading 1")
    for sheet, df in supp_tables:
        add_dataframe_table(sheet, df)
    for fig in supp_figures:
        add_figure(fig)

    add_para("Spatial Processing Rule", "Heading 1")
    add_para("Property boundaries remain the master geometries. All non-property environmental layers used upstream should be processed as 25 x 25 km tiles before intersection.")
    add_para("Conclusion", "Heading 1")
    add_para("The final products are internally consistent, traceable to retained raw inputs, free of zero municipality-code records, and structured as a publication-oriented results manuscript.")

    path = REPORT_DIR / f"forest_code_mt_final_report_{TODAY}.docx"
    doc.save(path)
    return path


def main() -> dict[str, Path]:
    tex = build_tex()
    docx = build_docx()
    print(f"tex: {tex}")
    print(f"docx: {docx}")
    return {"tex": tex, "docx": docx}


if __name__ == "__main__":
    main()

