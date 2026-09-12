from __future__ import annotations

import importlib.util
import os
import textwrap
from pathlib import Path

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

import _00_paths as paths
import _10_kernel as kernel

SCRIPT_DIR = Path(__file__).resolve().parent
paths.set_env()
ROOT = paths.ROOT
OUT_DIR = paths.TABLES
FIG_DIR = paths.FIGURES
FIG_DIR.mkdir(parents=True, exist_ok=True)
CHART_SIZE_IN = (7.5, 7.5)
MAP_SIZE_IN = (7.5, 7.5)
FIG_FONT_SIZE = 10
INPUT_ORDER = kernel.SOURCE_ORDER
INPUT_LABELS = kernel.SOURCE_LABELS

Y_LABELS = {
    "n_properties": "Number of properties",
    "total_deficit_ha": "Total deficit (hectares)",
    "gta_total_deficit_ha": "GTA-linked total deficit (hectares)",
    "total_cattle_moved": "Cattle head",
}


def comma_number(value: float, _pos=None) -> str:
    if pd.isna(value):
        return ""
    value = float(value)
    if abs(value) >= 10:
        return f"{value:,.0f}"
    return f"{value:,.1f}".rstrip("0").rstrip(".")


def wrapped_labels(values: pd.Series | list[object], width: int = 14) -> list[str]:
    return [
        "\n".join(textwrap.wrap(str(v), width=width, break_long_words=False, break_on_hyphens=False)) or str(v)
        for v in values
    ]

def first_existing(candidates: list[Path]) -> Path:
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


MUN_GEOM = paths.MUNICIPALITIES
UF_GEOM = paths.STATES


spec = importlib.util.spec_from_file_location("fcm_gta", SCRIPT_DIR / "_30_build_gta.py")
if spec is None or spec.loader is None:
    raise RuntimeError("Could not load _30_build_gta.py")
fcm_gta = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fcm_gta)


def today_path(name: str, ext: str) -> Path:
    return OUT_DIR / f"{name}_{kernel.DATE}.{ext}"


def num(df: pd.DataFrame, col: str) -> pd.Series:
    if col not in df.columns:
        return pd.Series(0.0, index=df.index)
    return pd.to_numeric(df[col], errors="coerce").fillna(0.0)


def mun_code(df: pd.DataFrame) -> pd.Series:
    if "mun_geocodigo" in df.columns:
        raw = df["mun_geocodigo"]
    elif "MUNICIPIO_" in df.columns:
        raw = df["MUNICIPIO_"]
    else:
        raw = pd.Series(pd.NA, index=df.index)
    out = pd.to_numeric(raw, errors="coerce").astype("Int64").astype("string").str.zfill(7)
    return out.mask(out.str.fullmatch(r"0+").fillna(False))


def save_bar(df: pd.DataFrame, x: str, y: str, title: str, path: Path, hue: str | None = None) -> Path:
    df = df.copy()
    plt.rcParams.update({
        "font.size": FIG_FONT_SIZE,
        "axes.labelsize": FIG_FONT_SIZE,
        "xtick.labelsize": FIG_FONT_SIZE,
        "ytick.labelsize": FIG_FONT_SIZE,
        "legend.fontsize": FIG_FONT_SIZE,
        "legend.title_fontsize": FIG_FONT_SIZE,
    })
    if x == "input_file_type" and x in df.columns:
        df[x] = pd.Categorical(df[x].astype(str), INPUT_ORDER, ordered=True)
        df = df.sort_values(x, kind="stable")
        labels = df[x].astype(str).map(INPUT_LABELS).fillna(df[x].astype(str))
    else:
        labels = df[x].astype(str)
    plt.figure(figsize=CHART_SIZE_IN)
    if hue and hue in df.columns:
        if hue == "input_file_type":
            df[hue] = pd.Categorical(df[hue].astype(str), INPUT_ORDER, ordered=True)
        piv = df.pivot_table(index=x, columns=hue, values=y, aggfunc="sum", fill_value=0, observed=False)
        if hue == "input_file_type":
            piv = piv.reindex(columns=INPUT_ORDER)
            piv.columns = [INPUT_LABELS.get(str(c), str(c)) for c in piv.columns]
        ax = piv.plot(kind="bar", stacked=False, ax=plt.gca(), width=0.8)
        ax.set_xticklabels(wrapped_labels(piv.index.astype(str), width=14), rotation=25, ha="right", fontsize=9)
        plt.legend(
            title="Input source" if hue == "input_file_type" else hue,
            loc="center left",
            bbox_to_anchor=(1.02, 0.5),
            frameon=False,
            fontsize=FIG_FONT_SIZE,
            title_fontsize=FIG_FONT_SIZE,
        )
    else:
        plt.bar(labels, df[y], color="#2F6F4E")
        plt.gca().set_xticks(np.arange(len(labels)))
        plt.gca().set_xticklabels(wrapped_labels(labels, width=14), rotation=25, ha="right", fontsize=9)
    plt.ylabel(Y_LABELS.get(y, y.replace("_", " ")))
    plt.xlabel("")
    plt.gca().yaxis.set_major_formatter(FuncFormatter(comma_number))
    plt.gca().tick_params(axis="both", labelsize=FIG_FONT_SIZE)
    plt.grid(False)
    plt.tight_layout(rect=(0.0, 0.0, 0.74, 1.0) if hue and hue in df.columns else None)
    plt.savefig(path, dpi=300)
    plt.close()
    return path


def build_tables(fc: pd.DataFrame, joined: pd.DataFrame) -> dict[str, pd.DataFrame]:
    fc = fc.copy()
    fc["mun_geocodigo_norm"] = mun_code(fc)
    joined = joined.copy()
    joined["mun_geocodigo_norm"] = mun_code(joined)
    for d in (fc, joined):
        if "input_file_type" in d.columns:
            d["input_file_type"] = pd.Categorical(d["input_file_type"].astype(str), INPUT_ORDER, ordered=True)
        d["area_ha_car"] = num(d, "area_ha_car")
        d["rl_adj_deficit_ha"] = num(d, "rl_adj_deficit_ha")
        d["rl_gross_deficit_ha"] = num(d, "rl_gross_deficit_ha")
        d["app_gross_deficit_ha"] = num(d, "app_gross_deficit_ha")
        d["calc_deficit_total_ha"] = num(d, "calc_deficit_total_ha")
        d["total_cattle_moved"] = num(d, "total_cattle_moved")

    mun_fc = (
        fc.groupby("mun_geocodigo_norm", dropna=False)
        .agg(
            n_properties=("priority_key", "size"),
            total_area_ha=("area_ha_car", "sum"),
            rl_adjusted_deficit_ha=("rl_adj_deficit_ha", "sum"),
            rl_gross_deficit_ha=("rl_gross_deficit_ha", "sum"),
            app_gross_deficit_ha=("app_gross_deficit_ha", "sum"),
            total_deficit_ha=("calc_deficit_total_ha", "sum"),
        )
        .reset_index()
    )

    mun_gta = (
        joined.groupby("mun_geocodigo_norm", dropna=False)
        .agg(
            n_gta_properties=("priority_key", "size"),
            total_cattle_head=("total_cattle_moved", "sum"),
            gta_total_deficit_ha=("calc_deficit_total_ha", "sum"),
        )
        .reset_index()
    )

    source_chart = (
        fc.groupby("input_file_type", observed=False)
        .agg(n_properties=("priority_key", "size"), total_deficit_ha=("calc_deficit_total_ha", "sum"))
        .reset_index()
    )
    supplier_chart = (
        joined.groupby(["input_file_type", "supplier_type"], observed=False)
        .agg(n_properties=("priority_key", "size"), total_deficit_ha=("calc_deficit_total_ha", "sum"))
        .reset_index()
    )
    size_chart = (
        fc.groupby(["input_file_type", "size_class"], observed=False, dropna=False)
        .agg(n_properties=("priority_key", "size"), total_deficit_ha=("calc_deficit_total_ha", "sum"))
        .reset_index()
    )
    source_chart = source_chart.sort_values("input_file_type", kind="stable")
    supplier_chart = supplier_chart.sort_values(["input_file_type", "supplier_type"], kind="stable")
    size_chart = size_chart.sort_values(["input_file_type", "size_class"], kind="stable")
    return {
        "municipal_fc": mun_fc,
        "municipal_gta": mun_gta,
        "chart_source": source_chart,
        "chart_supplier": supplier_chart,
        "chart_size": size_chart,
    }


def make_maps(mun_tbl: pd.DataFrame) -> list[Path]:
    paths: list[Path] = []
    if not MUN_GEOM.exists():
        return paths
    mun = gpd.read_parquet(MUN_GEOM)
    mun = mun[mun["geocodigo"].astype(str).str.startswith("51")].copy()
    mt = gpd.read_parquet(UF_GEOM)
    mt = mt[mt["sigla"].astype(str).eq("MT")].copy()
    gdf = mun.merge(mun_tbl, left_on="geocodigo", right_on="mun_geocodigo_norm", how="left")

    for col, title, cmap in [
        ("total_deficit_ha", "Total Forest Code deficit by municipality", "YlOrRd"),
        ("rl_adjusted_deficit_ha", "Adjusted Legal Reserve deficit by municipality", "YlGn"),
        ("app_gross_deficit_ha", "APP gross deficit by municipality", "Blues"),
    ]:
        gdf[col] = pd.to_numeric(gdf[col], errors="coerce").fillna(0)
        fig, ax = plt.subplots(figsize=MAP_SIZE_IN)
        gdf.plot(
            column=col,
            ax=ax,
            cmap=cmap,
            linewidth=0.15,
            edgecolor="white",
            legend=True,
            legend_kwds={"label": title + " (hectares)", "format": FuncFormatter(comma_number)},
        )
        if not mt.empty:
            mt.boundary.plot(ax=ax, color="#333333", linewidth=0.8)
        if len(fig.axes) > 1:
            cax = fig.axes[-1]
            cax.tick_params(labelsize=FIG_FONT_SIZE)
            cax.yaxis.label.set_size(FIG_FONT_SIZE)
        ax.set_axis_off()
        path = FIG_DIR / f"map_{col}.png"
        plt.tight_layout()
        plt.savefig(path, dpi=300)
        plt.close(fig)
        paths.append(path)
    return paths


def append_to_workbook(xlsx: Path, tables: dict[str, pd.DataFrame], figure_paths: list[Path]) -> None:
    wb = load_workbook(xlsx)
    for sheet in ["figures_index", "municipal_fc", "municipal_gta"]:
        if sheet in wb.sheetnames:
            del wb[sheet]

    ws = wb.create_sheet("figures_index")
    ws.append(["figure_file", "path"])
    for fig in figure_paths:
        ws.append([fig.name, str(fig)])
    row = 2
    for fig in figure_paths[:8]:
        img = XLImage(str(fig))
        img.width = 640
        img.height = 420
        ws.add_image(img, f"D{row}")
        row += 24

    for sheet, table in [("municipal_fc", tables["municipal_fc"]), ("municipal_gta", tables["municipal_gta"])]:
        ws_tbl = wb.create_sheet(sheet)
        ws_tbl.append(list(table.columns))
        for values in table.itertuples(index=False, name=None):
            ws_tbl.append(list(values))
    wb.save(xlsx)


def build_figures_and_maps() -> dict[str, Path]:
    final_outputs = fcm_gta.build_final_workbook()
    fc_path = today_path("forest_code_mt_priority_consolidated", "parquet")
    joined_path = today_path("forest_code_gta_final_mt", "parquet")
    fc = pd.read_parquet(fc_path)
    joined = pd.read_parquet(joined_path)
    tables = build_tables(fc, joined)

    figure_paths = [
        save_bar(tables["chart_source"], "input_file_type", "n_properties", "Properties by prioritized input", FIG_DIR / "chart_properties_by_input.png"),
        save_bar(tables["chart_source"], "input_file_type", "total_deficit_ha", "Total deficit by prioritized input", FIG_DIR / "chart_deficit_by_input.png"),
        save_bar(tables["chart_supplier"], "supplier_type", "n_properties", "GTA-matched properties by supplier type and input", FIG_DIR / "chart_supplier_by_input.png", hue="input_file_type"),
        save_bar(tables["chart_size"], "size_class", "total_deficit_ha", "Total deficit by size class and input", FIG_DIR / "chart_deficit_by_size.png", hue="input_file_type"),
    ]
    figure_paths.extend(make_maps(tables["municipal_fc"]))

    tables["municipal_fc"].to_csv(FIG_DIR / "municipal_fc_map_table.csv", index=False)
    tables["municipal_gta"].to_csv(FIG_DIR / "municipal_gta_map_table.csv", index=False)
    append_to_workbook(final_outputs["excel"], tables, figure_paths)

    print("Figures/maps written:")
    for fig in figure_paths:
        print(fig)
    print("Workbook updated:", final_outputs["excel"])
    return {"excel": final_outputs["excel"], "figures_dir": FIG_DIR}


if __name__ == "__main__":
    build_figures_and_maps()

