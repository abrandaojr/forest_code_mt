from __future__ import annotations

from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.path import Path as MplPath
from matplotlib.patches import PathPatch, Rectangle, FancyBboxPatch
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font
from PIL import Image as PILImage
import shutil

import _00_paths as paths
import _10_kernel as kernel


TABLE = paths.TABLES / f"forest_code_mt_priority_consolidated_with_secondary_{kernel.DATE}.parquet"
DELIVERY = paths.FIGURES
QA_TABLE = paths.TABLES / f"flow_figure_mass_balance_{kernel.DATE}.csv"
WORKBOOK = paths.TABLES / f"masson_style_final_tables_figures_{kernel.DATE}.xlsx"
DELIVERY_WORKBOOK = paths.TABLES / "masson_tables.xlsx"
GRAY = "#626E73"
LIGHT_GRAY = "#B5BDC1"
PALE_GRAY = "#E8ECEE"
INK = "#252525"
ORANGE = "#D97536"
FOREST = "#557F70"
OCHRE = "#D9A441"


def style() -> None:
    mpl.rcParams.update({
        "font.family": "Arial",
        "font.size": 12,
        "axes.labelsize": 12,
        "xtick.labelsize": 12,
        "ytick.labelsize": 12,
        "legend.fontsize": 12,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.facecolor": "white",
        "savefig.dpi": 300,
        "savefig.bbox": "tight",
        "savefig.pad_inches": 0.04,
    })


def fmt_million(value: float) -> str:
    return f"{value / 1_000_000:.2f} million ha"


def load_data() -> pd.DataFrame:
    columns = [
        "input_file_type", "rl_gross_deficit_ha", "rl_adj_deficit_ha",
        "rl_gross_deficit_forest_ha", "rl_gross_deficit_cerrado_ha",
        "rl_adj_deficit_forest_ha", "rl_adj_deficit_cerrado_ha",
        "rl_restore_ha", "rl_compensate_ha", "art67_small_prop",
        "art68_exempt_forest", "art68_exempt_cerrado", "mt_special_mun",
    ]
    return pd.read_parquet(TABLE, columns=columns)


def waterfall(df: pd.DataFrame) -> Path:
    gross = df["rl_gross_deficit_ha"].sum()
    reduction = df["rl_gross_deficit_ha"] - df["rl_adj_deficit_ha"]
    small = df["art67_small_prop"].fillna(False).astype(bool)
    art68_f = (~small) & df["art68_exempt_forest"].fillna(False).astype(bool)
    art68_c = (~small) & df["art68_exempt_cerrado"].fillna(False).astype(bool)
    # Component reductions follow the exact np.select precedence used by the model.
    art67 = reduction.where(small, 0).sum()
    forest_reduction = df["rl_gross_deficit_forest_ha"] - df["rl_adj_deficit_forest_ha"]
    cerrado_reduction = df["rl_gross_deficit_cerrado_ha"] - df["rl_adj_deficit_cerrado_ha"]
    art68 = forest_reduction.where(art68_f, 0).sum() + cerrado_reduction.where(art68_c, 0).sum()
    special = forest_reduction.where(
        ~small & ~art68_f & df["mt_special_mun"].fillna(False).astype(bool), 0
    ).sum()
    residual = reduction.sum() - art67 - art68 - special
    adjusted = df["rl_adj_deficit_ha"].sum()
    components = [art67, art68, special, residual]
    labels = ["Article 67", "Article 68", "MT special rule", "Other adjustment"]
    keep = [i for i, value in enumerate(components) if abs(value) > 0.5]
    components = [components[i] for i in keep]
    labels = [labels[i] for i in keep]

    fig, ax = plt.subplots(figsize=(10, 5.4))
    starts = [0.0]
    current = gross
    for value in components:
        starts.append(current - value)
        current -= value
    x = np.arange(len(components) + 2)
    ax.bar(0, gross, color=GRAY, width=0.62)
    for idx, (value, start) in enumerate(zip(components, starts[1:]), start=1):
        ax.bar(idx, value, bottom=start, color=LIGHT_GRAY, width=0.62)
        ax.plot([idx - 0.31, idx + 0.31], [start, start], color="#8B9498", lw=0.8)
    ax.bar(len(components) + 1, adjusted, color=ORANGE, width=0.62)
    for idx, value in enumerate([gross] + components + [adjusted]):
        y = gross if idx == 0 else (starts[idx] + value if idx <= len(components) else adjusted)
        prefix = "−" if 0 < idx <= len(components) else ""
        ax.text(idx, y + gross * 0.025, f"{prefix}{fmt_million(value)}", ha="center", va="bottom", fontsize=12, color=INK)
    ax.set_xticks(x, ["Gross LR deficit"] + labels + ["Adjusted LR deficit"])
    ax.set_ylabel("Legal Reserve deficit (million hectares)")
    ax.yaxis.set_major_formatter(lambda v, _: f"{v / 1_000_000:.1f}")
    ax.grid(axis="y", color="#DCE2E5", lw=0.7)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#777777")
    ax.tick_params(axis="x", rotation=0)
    fig.subplots_adjust(left=0.10, right=0.99, bottom=0.20, top=0.92)
    output = DELIVERY / "fig17_legal_reserve_deficit_adjustment_waterfall.png"
    fig.savefig(output)
    plt.close(fig)
    return output


def ribbon(ax, x0: float, x1: float, y0a: float, y0b: float, y1a: float, y1b: float, color: str, alpha: float = 0.60) -> None:
    control = (x1 - x0) * 0.48
    vertices = [
        (x0, y0a), (x0 + control, y0a), (x1 - control, y1a), (x1, y1a),
        (x1, y1b), (x1 - control, y1b), (x0 + control, y0b), (x0, y0b), (x0, y0a),
    ]
    codes = [MplPath.MOVETO, MplPath.CURVE4, MplPath.CURVE4, MplPath.CURVE4,
             MplPath.LINETO, MplPath.CURVE4, MplPath.CURVE4, MplPath.CURVE4, MplPath.CLOSEPOLY]
    ax.add_patch(PathPatch(MplPath(vertices, codes), facecolor=color, edgecolor="none", alpha=alpha))


def sankey(df: pd.DataFrame) -> Path:
    source_order = ["simcar_validado", "simcar_digital", "simcar_proxy"]
    source_labels = ["SIMCAR validated", "SIMCAR digital", "SIMCAR proxy"]
    grouped = df.groupby("input_file_type")[["rl_adj_deficit_ha", "rl_restore_ha", "rl_compensate_ha"]].sum().reindex(source_order).fillna(0)
    total = grouped["rl_adj_deficit_ha"].sum()
    scale = 0.72 / total
    gap = 0.035
    left_bottoms, cursor = [], 0.14
    for value in grouped["rl_adj_deficit_ha"]:
        left_bottoms.append(cursor)
        cursor += value * scale + gap
    middle_bottom = 0.14
    restore = grouped["rl_restore_ha"].sum()
    compensate = grouped["rl_compensate_ha"].sum()
    right_bottoms = [0.14, 0.14 + restore * scale + 0.10]

    fig, ax = plt.subplots(figsize=(10, 5.4))
    node_w = 0.025
    left_x, mid_x, right_x = 0.08, 0.50, 0.92
    mid_cursor = middle_bottom
    colors = ["#8A9499", "#657278", "#414C51"]
    for i, value in enumerate(grouped["rl_adj_deficit_ha"]):
        h = value * scale
        ribbon(ax, left_x + node_w, mid_x, left_bottoms[i], left_bottoms[i] + h, mid_cursor, mid_cursor + h, colors[i])
        ax.add_patch(Rectangle((left_x, left_bottoms[i]), node_w, h, color=colors[i]))
        ax.text(left_x - 0.015, left_bottoms[i] + h / 2, f"{source_labels[i]}\n{fmt_million(value)}", ha="right", va="center", fontsize=12, color=INK)
        mid_cursor += h
    ax.add_patch(Rectangle((mid_x, middle_bottom), node_w, total * scale, color=ORANGE))
    ax.text(mid_x + node_w / 2, middle_bottom + total * scale + 0.025, f"Adjusted LR deficit\n{fmt_million(total)}", ha="center", va="bottom", fontsize=12, color=INK)

    mid_out = middle_bottom
    for value, bottom, label, color in [
        (restore, right_bottoms[0], "Mandatory restoration", ORANGE),
        (compensate, right_bottoms[1], "Compensation pathway", GRAY),
    ]:
        h = value * scale
        ribbon(ax, mid_x + node_w, right_x, mid_out, mid_out + h, bottom, bottom + h, color)
        ax.add_patch(Rectangle((right_x, bottom), node_w, h, color=color))
        ax.text(right_x + node_w + 0.015, bottom + h / 2, f"{label}\n{fmt_million(value)}", ha="left", va="center", fontsize=12, color=INK)
        mid_out += h
    ax.text(0.5, 0.035, "Ribbon width represents adjusted Legal Reserve deficit area (hectares).", ha="center", va="center", fontsize=12, color="#555555")
    ax.set_xlim(0, 1.10)
    ax.set_ylim(0, 1)
    ax.axis("off")
    fig.subplots_adjust(left=0.03, right=0.97, bottom=0.02, top=0.98)
    output = DELIVERY / "fig18_legal_reserve_deficit_pathways_sankey.png"
    fig.savefig(output)
    plt.close(fig)

    checks = pd.DataFrame([
        {"check": "source_to_adjusted", "incoming_ha": total, "outgoing_ha": grouped["rl_adj_deficit_ha"].sum()},
        {"check": "adjusted_to_pathways", "incoming_ha": total, "outgoing_ha": restore + compensate},
    ])
    checks["difference_ha"] = checks["incoming_ha"] - checks["outgoing_ha"]
    checks.to_csv(QA_TABLE, index=False)
    if checks["difference_ha"].abs().max() > 1e-6:
        raise RuntimeError("Sankey mass-balance check failed")
    return output


def methodology() -> Path:
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    boxes = [
        (0.02, 0.70, 0.25, 0.20, "CAR property geometries\nvalidated · digital · proxy", LIGHT_GRAY),
        (0.02, 0.40, 0.25, 0.20, "RADAMBrasil classification\nforest · Cerrado", OCHRE),
        (0.02, 0.10, 0.25, 0.20, "Native vegetation and\nForest Code spatial inputs", FOREST),
        (0.36, 0.35, 0.28, 0.30, "Equal-area spatial intersection\ngeometry repair\narea measured in hectares\ncoverage and key audits", GRAY),
        (0.73, 0.58, 0.25, 0.20, "Property-level classification\nrequirements and\nexisting vegetation", LIGHT_GRAY),
        (0.73, 0.22, 0.25, 0.20, "Compliance outcomes\ndeficit · surplus\nrestoration · compensation", ORANGE),
    ]
    for x, y, w, h, label, color in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.012,rounding_size=0.012", facecolor=color, edgecolor="#555555", lw=0.9))
        text_color = "white" if color in {GRAY, FOREST, ORANGE} else INK
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=12, color=text_color, linespacing=1.25)
    for start, end in [
        ((0.27, 0.80), (0.36, 0.57)), ((0.27, 0.50), (0.36, 0.50)), ((0.27, 0.20), (0.36, 0.43)),
        ((0.64, 0.50), (0.73, 0.68)), ((0.855, 0.58), (0.855, 0.42)),
    ]:
        ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle": "-|>", "color": "#555555", "lw": 1.4, "shrinkA": 4, "shrinkB": 4})
    ax.text(0.50, 0.17, "Outputs are consolidated by priority\nwithout double-counting properties.", ha="center", va="center", fontsize=12, color="#555555")
    output = DELIVERY / "fig19_spatial_data_integration_workflow.png"
    fig.savefig(output)
    plt.close(fig)
    return output


def refresh_workbook_figure_index() -> None:
    figures = sorted(DELIVERY.glob("*.png"))
    workbook = load_workbook(WORKBOOK)
    if "Figures_Index" in workbook.sheetnames:
        del workbook["Figures_Index"]
    sheet = workbook.create_sheet("Figures_Index")
    sheet.append(["figure", "path"])
    sheet.column_dimensions["A"].width = 58
    sheet.column_dimensions["B"].width = 95
    row = 2
    for figure in figures:
        sheet.append([figure.name, str(figure)])
        with PILImage.open(figure) as source_image:
            aspect = source_image.height / source_image.width
        image = XLImage(str(figure))
        image.width = 720
        image.height = int(720 * aspect)
        sheet.add_image(image, f"D{row}")
        row += max(24, int(image.height / 20) + 2)
    for row_cells in sheet.iter_rows(min_col=1, max_col=2):
        for cell in row_cells:
            cell.font = Font(name="Arial", size=12, bold=cell.row == 1)
    workbook.save(WORKBOOK)
    shutil.copy2(WORKBOOK, DELIVERY_WORKBOOK)


def main() -> None:
    style()
    DELIVERY.mkdir(parents=True, exist_ok=True)
    df = load_data()
    outputs = [waterfall(df), sankey(df), methodology()]
    refresh_workbook_figure_index()
    print("Generated:")
    for output in outputs:
        print(output)
    print(f"Mass-balance audit: {QA_TABLE}")


if __name__ == "__main__":
    main()
