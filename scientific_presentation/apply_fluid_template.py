from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "qa" / "presentation" / "fluid_template"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

PINK = "#EF3F6B"
YELLOW = "#FFE266"
TEAL = "#10B7B5"
PURPLE = "#40288A"
SILVER = "#EEEEEF"


def silver_background(variant: int) -> Path:
    width, height = 1920, 1080
    image = Image.new("RGB", (width, height), SILVER)
    pixels = image.load()
    for y in range(height):
        for x in range(width):
            glow = int(15 * (1 - abs(x / width - 0.42))) + int(7 * (1 - y / height))
            value = max(224, min(250, 232 + glow))
            pixels[x, y] = (value, value, min(255, value + 2))
    draw = ImageDraw.Draw(image)
    if variant == 0:
        draw.ellipse((-260, 900, 260, 1360), fill=YELLOW)
        draw.ellipse((1590, -260, 2170, 280), fill=TEAL)
        draw.ellipse((1760, 920, 2100, 1240), fill=PINK)
    elif variant == 1:
        draw.ellipse((1650, 900, 2170, 1370), fill=PINK)
        draw.ellipse((1660, -220, 2050, 170), fill=YELLOW)
    else:
        draw.ellipse((-260, 910, 260, 1360), fill=YELLOW)
        draw.ellipse((1610, -250, 2160, 270), fill=PINK)
    target = ASSET_DIR / f"fluid-background-{variant + 1}.png"
    image.save(target)
    return target


def clamp_font_sizes(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    current = run.font.size.pt if run.font.size else None
                    run.font.name = "Arial"
                    if current is None:
                        run.font.size = Pt(20)
                    else:
                        run.font.size = Pt(max(20, min(40, current)))


def apply_backgrounds(pptx: Path) -> None:
    backgrounds = [silver_background(i) for i in range(3)]
    prs = Presentation(pptx)
    for index, slide in enumerate(prs.slides):
        picture = slide.shapes.add_picture(
            str(backgrounds[index % len(backgrounds)]),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        tree = slide.shapes._spTree
        tree.remove(picture._element)
        tree.insert(2, picture._element)
    clamp_font_sizes(prs)
    prs.save(pptx)


if __name__ == "__main__":
    apply_backgrounds(
        ROOT
        / "deliverables"
        / "04_presentation"
        / "01_FOREST_CODE_TWO_PROPERTY_WORKED_EXAMPLES_20260930.pptx"
    )
